"""Add and fill Supply Method on CL, Plain Database, and Packs.

  python scripts/fill_supply_method.py --dry-run
  python scripts/fill_supply_method.py --target cl --dry-run
"""
from __future__ import annotations
import argparse
import csv
import shutil
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
from openpyxl import load_workbook
from shared import paths as wh
from shared.areeb_taxonomy import cell
from shared.supply_method import (
    COL,
    classify_cl_row,
    classify_packs_row,
    classify_plain_row,
)
PLAIN_SHEET = "Sheet1"
PACKS_SHEET = "01-Database"
def _print_counts(title: str, path: Path, stats: dict[str, int], samples: list[str]) -> None:
    print(f"{title}: {path}")
    print(f"  rows={stats['rows']:,}")
    for key in sorted(k for k in stats if k != "rows"):
        print(f"  {key}: {stats[key]:,}")
    for line in samples:
        print(f"  e.g. {line}")
def fill_cl(*, dry_run: bool) -> dict[str, int]:
    path = wh.cl_csv_path()
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        headers = list(reader.fieldnames or [])
        rows = [{k: cell(v) for k, v in row.items()} for row in reader]
    if COL not in headers:
        headers.append(COL)
    stats: dict[str, int] = defaultdict(int)
    samples: list[str] = []
    for row in rows:
        value = classify_cl_row(row)
        prev = cell(row.get(COL))
        stats["rows"] += 1
        stats[value] += 1
        if prev != value:
            stats["would_write" if dry_run else "wrote"] += 1
        row[COL] = value
        if len(samples) < 6:
            samples.append(
                f"{row.get('Custom Label')} ga={row.get('Gender Apparel')!r} -> {value}"
            )
    _print_counts("Custom Label", path, dict(stats), samples)
    if dry_run:
        return dict(stats)
    bak = backup_file(path, wh.cl_backups_dir())
    print(f"  backup {bak}")
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    print("  wrote", path)
    return dict(stats)
def _ensure_xlsx_col(ws, headers: list[str], idx: dict[str, int]) -> dict[str, int]:
    if COL in idx:
        return idx
    col_i = len(headers) + 1
    ws.cell(1, col_i).value = COL
    headers.append(COL)
    idx[COL] = col_i - 1
    return idx
def fill_plain(*, dry_run: bool) -> dict[str, int]:
    path = wh.plain_database_path()
    wb = load_workbook(path, read_only=dry_run, data_only=dry_run)
    ws = wb[PLAIN_SHEET]
    headers = [cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers)
    stats: dict[str, int] = defaultdict(int)
    samples: list[str] = []
    if dry_run:
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = {h: values[i] if i < len(values) else "" for i, h in enumerate(headers)}
            value = classify_plain_row(row)
            stats["rows"] += 1
            stats[value] += 1
            if cell(row.get(COL)) != value:
                stats["would_write"] += 1
            if len(samples) < 6:
                samples.append(f"SKU={row.get('SKU')} {row.get('Brand')} -> {value}")
        wb.close()
        _print_counts("Plain Database", path, dict(stats), samples)
        return dict(stats)
    idx = _ensure_xlsx_col(ws, headers, idx)
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(r, i + 1).value for h, i in idx.items()}
        value = classify_plain_row(row)
        prev = cell(row.get(COL))
        stats["rows"] += 1
        stats[value] += 1
        if prev != value:
            stats["wrote"] += 1
            ws.cell(r, idx[COL] + 1).value = value
        if len(samples) < 6:
            samples.append(f"SKU={row.get('SKU')} {row.get('Brand')} -> {value}")
    _print_counts("Plain Database", path, dict(stats), samples)
    bak = backup_file(path, wh.plain_database_archive_dir())
    print(f"  backup {bak}")
    wb.save(path)
    wb.close()
    print("  wrote", path)
    return dict(stats)
def fill_packs(*, dry_run: bool) -> dict[str, int]:
    path = wh.packs_database_path()
    wb = load_workbook(path, read_only=dry_run, data_only=dry_run)
    ws = wb[PACKS_SHEET]
    headers = [cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers)
    stats: dict[str, int] = defaultdict(int)
    samples: list[str] = []
    if dry_run:
        for values in ws.iter_rows(min_row=2, values_only=True):
            row = {h: values[i] if i < len(values) else "" for i, h in enumerate(headers)}
            value = classify_packs_row(row)
            stats["rows"] += 1
            stats[value] += 1
            if cell(row.get(COL)) != value:
                stats["would_write"] += 1
            if len(samples) < 6:
                samples.append(
                    f"{row.get('Channel Child SKU')} {row.get('Brand Name')} -> {value}"
                )
        wb.close()
        _print_counts("Packs Database", path, dict(stats), samples)
        return dict(stats)
    idx = _ensure_xlsx_col(ws, headers, idx)
    for r in range(2, ws.max_row + 1):
        row = {h: ws.cell(r, i + 1).value for h, i in idx.items()}
        value = classify_packs_row(row)
        prev = cell(row.get(COL))
        stats["rows"] += 1
        stats[value] += 1
        if prev != value:
            stats["wrote"] += 1
            ws.cell(r, idx[COL] + 1).value = value
        if len(samples) < 6:
            samples.append(f"{row.get('Channel Child SKU')} {row.get('Brand Name')} -> {value}")
    _print_counts("Packs Database", path, dict(stats), samples)
    bak = backup_file(path, wh.packs_database_archive_dir())
    print(f"  backup {bak}")
    wb.save(path)
    wb.close()
    print("  wrote", path)
    return dict(stats)
