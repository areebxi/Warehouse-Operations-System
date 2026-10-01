"""Simple CEO progress workbook. Re-run after locks change."""
from __future__ import annotations
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
NAVY = "1F4E79"
WHITE = "FFFFFF"
GREEN = "C6EFCE"
GREEN_F = "006100"
YELLOW = "FFF2CC"
YELLOW_F = "7F6000"
ORANGE = "FCE4D6"
ORANGE_F = "C65911"
BLUE = "DDEBF7"
GREY = "F2F2F2"
C = Alignment(wrap_text=True, vertical="center", horizontal="center")
from write_order_grouping_progress_style import (
    f,
    put,
    widths,
    L,
    LINE,
)
def sheet_board(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "Board"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 26
    put(ws, 1, 1, "Order grouping  ·  " + date.today().isoformat(), bg=NAVY, fg=WHITE, bold=True, size=16, align=L)
    ws.merge_cells("A1:L1")
    put(ws, 2, 1, "Green = locked     Yellow = off (v1)     Graph values = examples except Hashim #038 Areeb pick-lists (Title Case; type has no gender; style is a product name)     1 logic ✓  2 catalog fill ✓  3 sorter ✓  4 Input write 2026-09-11", bg=BLUE, bold=True, align=L)
    ws.merge_cells("A2:L2")

    # Process name
    put(ws, 4, 1, "Process name", bg=NAVY, fg=WHITE, bold=True, align=C)
    slots = [
        ("PLAIN/PRINTED", GREEN, "product-finish"),
        ("1 / 2", GREEN, "prime / non-prime"),
        ("WAREHOUSE STOCK", GREEN, "or SUPPLY ON DEMAND / IN HOUSE MANUFACTURE"),
        ("P / R", GREEN, "personalised / ready made"),
        ("S1", GREEN, "testing: every --run is 1st Shift; production later S2 / S3"),
        ("1", GREEN, "priority — today first, then prime, then as made"),
    ]
    for i, (name, bg, meaning) in enumerate(slots):
        col = 2 + i
        put(ws, 4, col, name, bg=bg, bold=True, size=12, align=C)
        put(ws, 5, col, meaning, bg=GREY, size=9, align=C)
    put(ws, 5, 1, "printed ex.", bg=GREY, size=9, align=C)
    ws.row_dimensions[5].height = 36

    put(ws, 6, 1, "plain name", bg=GREY, size=9, align=C)
    put(ws, 6, 2, "Fixed batches: B40/B80/B100/B1050/B3700/…. Leftover: B1/B2/… (skip reserved). Filename keeps full 6 fields. Packing PIN uses B100-S1-1 Item 1 only. Shift folder stays 1st/2nd/3rd. Amazon stays in own. MAS Clothing → fawad (match only).", bg=GREEN, fg=GREEN_F, size=9, align=L)
    ws.merge_cells("B6:J6")

    put(ws, 8, 1, "RESEND", bg=YELLOW, fg=YELLOW_F, bold=True, align=C)
    put(ws, 8, 2, "filename is only:  RESEND     Exact tag 1014-ALL-RESEND only (not 1015 / 1016 / 1017). Wins even if ship-by is blank.", bg=YELLOW, align=L)
    ws.merge_cells("B8:J8")
    put(ws, 9, 1, "UNMATCHED", bg=ORANGE, fg=ORANGE_F, bold=True, align=C)
    put(ws, 9, 2, "filename is only:  UNMATCHED     No definite match. Floor handles this run; then investigate and tighten so the next run matches. Same order stays together.", bg=ORANGE, align=L)
    ws.merge_cells("B9:J9")

    put(ws, 11, 1, "This shift", bg=NAVY, fg=WHITE, bold=True, align=C)
    ws.merge_cells("A11:F11")
    intake = [
        "1. ShipStation\nawaiting_shipment only",
        "2. Skip\npost-order-designs",
        "3. Resend tag?\n→ file “RESEND”",
        "4. Blank ship-by\n→ today",
        "5. Today/overdue first\nthen future to fill cap",
        "6. Folders 1st/2nd/3rd\n300/100/100 LINES",
    ]
    for i, t in enumerate(intake):
        put(ws, 12, 1 + i, t, bg=GREEN, fg=GREEN_F, bold=True, size=10, align=C)
        ws.merge_cells(start_row=12, start_column=1 + i, end_row=13, end_column=1 + i)
    ws.row_dimensions[12].height = 22
    ws.row_dimensions[13].height = 28

    # Splits
    put(ws, 15, 1, "New process file when flag is 1", bg=NAVY, fg=WHITE, bold=True, align=L)
    ws.merge_cells("A15:J15")
    headers = ["Split", "Plain", "Printed", "Where", "v1"]
    for i, h in enumerate(headers):
        put(ws, 16, 1 + i, h, bg=NAVY, fg=WHITE, bold=True, align=C)
    splits = [
        ("shift", "on", "on", "Testing: filename field 1 S1 (after optional B-batch prefix) and Input 1st Shift every --run. Hashim #037: order volume > 300 splits today vs later dates; do not dump into Shift 2/3. Blank ship-by → today (in awaiting_shipment → pull today; 2026-09-24). Production later: nth run = nth shift", "on"),
        ("plain / printed", "on", "on", "Finish gate: SKU plain/plainlg → plain; else CL hit (after 1st dash, or whole SKU if no dash) → printed; else Plain Database till-last-dash or Packs whole → plain; else unmatched. Filename PLAIN or PRINTED", "on"),
        ("store / channel", "on", "on", "MAS Clothing → fawad match (not a filename slot); Amazon stays in own", "on"),
        ("design group", "off", "off", "per scenario, Design ID lists", "off"),
        ("prime", "on", "on", "tag Amazon Prime Order (not CL Amazon Prime)", "on"),
        ("printing method", "off", "on", "CL Printing Type — DTF default; Sublimation = mugs", "on"),
        ("ready-made / customised", "off", "on", "CL Customise. Mixed P vs R in one order: majority printed units; tie → readymade (not unmatched)", "on"),
        ("customisation type", "off", "off", "CL Customisation Type (column empty; v1 slot x)", "off"),
        ("print size", "off", "off", "CL Print Size 1", "off"),
        ("print position", "off", "off", "CL Print Positions", "off"),
        ("supply method", "on", "on", "CL Warehouse Stock = FOTL tees in locked colour lists (Mens/Womens/Kids) plus Kids C800T/C8030T body colours. Filled 2026-09-23 (63,399 warehouse). Plain/Packs Warehouse Stock = all FOTL tees. In House = iron-on/sticker (printed). Else Supplier On Demand. Plain never in-house. Mixed supply-method → Supplier On Demand.", "on"),
        ("supplier", "on", "stock off / on-demand on", "Supplier Name: BTC Activewear / Uneek Clothing / Absolute Apparels (babysuits C800T C8020T C8030T only)", "on"),
        ("package type", "off", "off", "packing PDF only (Large Letter / Parcel). Packs vs Gildan = 30-chain", "off"),
    ]

    def flag_bg(v: str) -> str:
        if v == "on":
            return GREEN
        if v == "off":
            return YELLOW
        if v.startswith("on"):
            return ORANGE
        return GREY

    for r, row in enumerate(splits, start=17):
        put(ws, r, 1, row[0], bold=True, align=L)
        put(ws, r, 2, row[1], bg=flag_bg(row[1]), align=C)
        put(ws, r, 3, row[2], bg=flag_bg(row[2]), align=C)
        put(ws, r, 4, row[3], align=L)
        put(ws, r, 5, row[4], bg=flag_bg(row[4]), bold=True, align=C)
    put(ws, 30, 1, "Fixed batches first (B10…B8050, row order = priority), then plain B2000/B2100/B2200/B2500/B2600/… and printed leftover B1/B2/… skipping reserved. Hashim #037: orders > 300 split today vs later. Testing: --run rewrites 1st Shift. Packing PIN: B100-S1-1 Item 1 (batch+shift only).", bg=GREEN, fg=GREEN_F, align=L)
    ws.merge_cells("A30:E30")

    # 30 chain
    put(ws, 32, 1, "Then, if a bucket has 30+ units → another process file", bg=NAVY, fg=WHITE, bold=True, align=L)
    ws.merge_cells("A32:I32")
    chain = [
        ("Category (Areeb)", GREEN, "filled"),
        ("Product Type (Areeb)", GREEN, "filled"),
        ("Product Style (Areeb)", GREEN, "filled"),
        ("Department (Areeb)", GREEN, "30 both"),
        ("Brand / Brand Name", GREEN, "30 both"),
        ("Size / Pack Size", GREEN, "30 both"),
        ("Colour", GREEN, "have; packs later"),
    ]
    for i, (name, bg, note) in enumerate(chain):
        put(ws, 33, 1 + i, name, bg=bg, bold=True, align=C)
        put(ws, 34, 1 + i, note, bg=bg, size=9, align=C)

    put(ws, 35, 1, "Fixed-batch FOTL / iron-on families skip the 30-chain (already their own file). Leftover Graph 30-chain (plain AND printed): every step 30. Packs skip colour.", bg=GREY, size=10, align=L)
    ws.merge_cells("A35:I35")

    put(ws, 36, 1, "If a branch is not drawn all the way to colour, it still uses this chain (blank = save writing, not skip).", bg=YELLOW, fg=YELLOW_F, bold=True, align=L)
    ws.merge_cells("A36:I36")
    put(ws, 37, 1, "Printed in-house (iron-on or sticker) uses the same 30-chain tail. Blank Brand/Colour on flag 30 stays in the parent (not unmatched). Plain cannot be in-house. Mixed plain+printed → whole order printed.", bg=YELLOW, align=L)
    ws.merge_cells("A37:I37")

    put(ws, 39, 1, "On each packing line — batch, then process number, then item number", bg=NAVY, fg=WHITE, bold=True, align=L)
    ws.merge_cells("A39:J39")
    put(ws, 40, 1, "B80-S1-1 Item 1     then     B80-S1-1 Item 2     then     B80-S1-2 Item 1", bg=GREEN, fg=GREEN_F, bold=True, align=C)
    ws.merge_cells("A40:J40")
    put(ws, 41, 1, "PIN uses batch+shift only (B80-S1), not the full filename. Process number = colour 3+ then 50 units. Item number = line. Same order keeps the same process number. Packs skip colour groups.", bg=GREY, align=L)
    ws.merge_cells("A41:J41")

    widths(ws, {**{chr(65 + i): 16 for i in range(12)}, "A": 26, "B": 18, "C": 28, "D": 42, "E": 22})
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["C"].width = 28
    ws.column_dimensions["D"].width = 44
    ws.column_dimensions["E"].width = 18
    for col in "FGHIJ":
        ws.column_dimensions[col].width = 16
    ws.freeze_panes = "A3"
    ws.print_area = "A1:J41"
    ws.page_setup.horizontalCentered = True
    ws.oddHeader.left.text = "Order grouping progress"
