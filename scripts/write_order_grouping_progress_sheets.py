"""Simple CEO progress workbook. Re-run after locks change."""
from __future__ import annotations
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
from write_order_grouping_progress_style import (
    f,
    put,
    widths,
    L,
    LINE,
)
from write_order_grouping_progress_board import (
    sheet_board,
    BLUE,
    C,
    GREEN,
    GREEN_F,
    GREY,
    NAVY,
    ORANGE,
    ORANGE_F,
    WHITE,
    YELLOW,
    YELLOW_F,
)
def sheet_catalog(wb: Workbook) -> None:
    ws = wb.create_sheet("Catalog")
    ws.page_setup.orientation = "landscape"
    ws.sheet_view.showGridLines = False
    put(ws, 1, 1, "Step 2 — catalog fill done (2026-09-08/09). CL Areeb squeezed refill 2026-09-16. Sorter reads; it does not overwrite.", bg=NAVY, fg=WHITE, bold=True, size=16, align=L)
    ws.merge_cells("A1:D1")
    put(ws, 2, 1, "30-chain = Areeb columns. Do not overwrite PE Department. Packs colour later. Plain/Packs Areeb stay supplier copy. Backup before any future fill.", bg=YELLOW, fg=YELLOW_F, bold=True, align=L)
    ws.merge_cells("A2:D2")
    for i, h in enumerate(["Column", "Do this", "Used for", "v1"]):
        put(ws, 4, 1 + i, h, bg=NAVY, fg=WHITE, bold=True, align=C)
    rows = [
        ("Customise (CL)", "have (Yes / blank)", "ready-made vs customised", GREEN, "on"),
        ("Printing Type (CL)", "FILLED — DTF / mug Sublimation", "printed printing method", GREEN, "on"),
        ("Supply Method (all 3)", "FILLED — FOTL tees / iron-on+sticker / on-demand", "Warehouse Stock, In House Manufacture, Supplier On Demand", GREEN, "on"),
        ("Package Type / Package", "CL 34%; Plain 100%; display", "packing PDF; grouping off", YELLOW, "off"),
        ("Supplier Name", "FILLED — BTC Activewear / Uneek Clothing / Absolute babysuits", "supplier split when flag 1", GREEN, "on"),
        ("Category (Areeb)", "FILLED; pick-list #038 Title Case", "≥30 new file", GREEN, "on"),
        ("Product Type (Areeb)", "FILLED; pick-list #038 no gender", "≥30 new file", GREEN, "on"),
        ("Product Style (Areeb)", "FILLED; named product not code", "≥30 new file", GREEN, "on"),
        ("Department (Areeb)", "FILLED; CL = gender only (on list)", "plain 30 / printed 1", GREEN, "on"),
        ("Brand, Size, Colour", "have (Packs: Brand Name + Pack Size; colour later)", "30+ files; colour 3+ qty → -N (not packs). Same 30 on printed.", GREEN, "on"),
        ("Print Size 1", "leave; fill later", "print size split (v2)", YELLOW, "off"),
        ("Print Positions", "have; split later", "print position split (v2)", YELLOW, "off"),
        ("Amazon Prime (CL)", "leave — not grouping", "PO / inventory only", YELLOW, "not grouping"),
        ("customisation type", "CL column empty; v1 off", "name slot x in v1", YELLOW, "off"),
        ("Plain Database.xlsx", "grouping run is READ", "till last dash → SKU", GREEN, "match"),
        ("Packs Database.xlsx", "grouping run is READ", "whole SKU → Channel Child SKU", GREEN, "match"),
    ]
    for r, (col, do, used, bg, v1) in enumerate(rows, start=5):
        put(ws, r, 1, col, bold=True)
        put(ws, r, 2, do, bg=bg, align=C)
        put(ws, r, 3, used)
        put(ws, r, 4, v1, bg=bg, bold=True, align=C)
        ws.row_dimensions[r].height = 22
    put(ws, 22, 1, "Step 3: Order Grouping Sorter. SKU keys: CL after 1st dash (no dash → whole) | Plain till last dash | Packs whole. Attrs: Packs→CL→Plain. Packs colour later. Mixed supply → on-demand.", bg=GREEN, fg=GREEN_F, align=L)
    ws.merge_cells("A22:D22")
    widths(ws, {"A": 24, "B": 32, "C": 40, "D": 16})
def sheet_design(wb: Workbook) -> None:
    ws = wb.create_sheet("Design")
    ws.sheet_view.showGridLines = False
    put(ws, 1, 1, "Design grouping — off in v1. Different list per scenario.", bg=NAVY, fg=WHITE, bold=True, size=16, align=L)
    ws.merge_cells("A1:E1")
    put(ws, 2, 1, "When you turn it on, process name slot = group-01, group-02…  Fill a scenario before a run.", bg=BLUE, align=L)
    ws.merge_cells("A2:E2")
    for i, h in enumerate(["Scenario", "On (0/1)", "Group", "Logo / Design IDs", "Notes"]):
        put(ws, 4, 1 + i, h, bg=NAVY, fg=WHITE, bold=True, align=C)
    put(ws, 5, 1, "default (now)")
    put(ws, 5, 2, "0", bg=YELLOW, align=C)
    put(ws, 5, 3, "x", align=C)
    put(ws, 5, 4, "—", align=C)
    put(ws, 5, 5, "v1")
    put(ws, 6, 1, "CEO example")
    put(ws, 6, 2, "1", align=C)
    put(ws, 6, 3, "group-01", align=C)
    put(ws, 6, 4, "123LG, 245LG")
    put(ws, 6, 5, "sample only")
    for r in range(7, 16):
        for c in range(1, 6):
            put(ws, r, c, "")
    widths(ws, {"A": 20, "B": 12, "C": 14, "D": 36, "E": 22})
