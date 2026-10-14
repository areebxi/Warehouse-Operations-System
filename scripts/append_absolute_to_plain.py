"""Append Absolute Product Data stock codes onto Plain Database.csv.

Existing SKUs are left as they are. Rows that exist only on the old xlsx
are copied across first so the CSV stays the full plain catalog.

  python scripts/append_absolute_to_plain.py
"""

from __future__ import annotations

import csv
import shutil
import sys
from datetime import datetime
from pathlib import Path

_ROOT = Path(__file__).resolve().parents[1]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from shared import paths as wh  # noqa: E402
from shared.absolute_catalog import grouping_row, load_absolute_rows  # noqa: E402

COLUMNS = [
    "SKU",
    "Product Code",
    "Brand",
    "Colour",
    "Size",
    "Description",
    "Product_Image_URL",
    "Brand_Image_URL",
    "Package",
    "Category (Areeb)",
    "Product Type (Areeb)",
    "Product Style (Areeb)",
    "Department (Areeb)",
    "Supply Method",
    "Supplier Name",
]


def _cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _read_csv(path: Path) -> list[dict[str, str]]:
    last: Exception | None = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with path.open(encoding=enc, newline="") as handle:
                return [{k: _cell(v) for k, v in row.items() if k} for row in csv.DictReader(handle)]
        except UnicodeDecodeError as exc:
            last = exc
    raise last or RuntimeError(f"Could not decode {path}")


def _xlsx_only(csv_skus: set[str]) -> list[dict[str, str]]:
    xlsx = wh.plain_database_path().with_name("Plain Database.xlsx")
    if not xlsx.is_file():
        return []
    from openpyxl import load_workbook

    wb = load_workbook(xlsx, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        headers = [_cell(h) for h in next(it)]
        extra: list[dict[str, str]] = []
        for raw in it:
            row = {
                headers[i]: _cell(raw[i] if i < len(raw) else "")
                for i in range(len(headers))
                if headers[i]
            }
            sku = row.get("SKU", "")
            if sku and sku.casefold() not in csv_skus:
                extra.append(row)
                csv_skus.add(sku.casefold())
        return extra
    finally:
        wb.close()


def _absolute_rows(seen: set[str]) -> list[dict[str, str]]:
    out: list[dict[str, str]] = []
    for raw in load_absolute_rows():
        grouped = grouping_row(raw)
        sku = grouped.get("SKU", "")
        if not sku or sku.casefold() in seen:
            continue
        seen.add(sku.casefold())
        row = {col: "" for col in COLUMNS}
        row.update({k: grouped.get(k, "") for k in COLUMNS if k in grouped})
        row["Package"] = "Small Parcel"
        row["Product_Image_URL"] = ""
        row["Brand_Image_URL"] = ""
        out.append(row)
    return out


def main() -> None:
    path = wh.plain_database_path().with_name("Plain Database.csv")
    rows = _read_csv(path)
    seen = {_cell(r.get("SKU")).casefold() for r in rows if _cell(r.get("SKU"))}
    from_xlsx = _xlsx_only(seen)
    added = _absolute_rows(seen)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    archive = wh.plain_database_archive_dir()
    archive.mkdir(parents=True, exist_ok=True)
    backup = archive / f"Plain Database.bak_{stamp}.csv"
    shutil.copy2(path, backup)
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=COLUMNS, extrasaction="ignore")
        writer.writeheader()
        for row in rows + from_xlsx + added:
            writer.writerow({col: _cell(row.get(col)) for col in COLUMNS})
    print(f"Kept {len(rows)} CSV rows, copied {len(from_xlsx)} from xlsx, added {len(added)} Absolute.")
    print(f"Backup: {backup}")


if __name__ == "__main__":
    main()
