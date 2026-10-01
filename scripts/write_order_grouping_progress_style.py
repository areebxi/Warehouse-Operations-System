"""Simple CEO progress workbook. Re-run after locks change."""
from __future__ import annotations
from datetime import date
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet
LINE = Border(
    left=Side(style="thin", color="BDD7EE"),
    right=Side(style="thin", color="BDD7EE"),
    top=Side(style="thin", color="BDD7EE"),
    bottom=Side(style="thin", color="BDD7EE"),
)
L = Alignment(wrap_text=True, vertical="center", horizontal="left")
def f(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)
def put(ws, r, c, val, *, bg=None, fg="000000", bold=False, size=11, align=L):
    cell = ws.cell(r, c, val)
    cell.font = Font(name="Calibri", bold=bold, color=fg, size=size)
    cell.alignment = align
    cell.border = LINE
    if bg:
        cell.fill = f(bg)
    return cell
def widths(ws: Worksheet, d: dict[str, float]) -> None:
    for k, v in d.items():
        ws.column_dimensions[k].width = v
