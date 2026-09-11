"""Fill Category/Product Type/Product Style/Department (Areeb) on CL, Plain, Packs.

  python scripts/fill_areeb_taxonomy.py --target plain --dry-run
  python scripts/fill_areeb_taxonomy.py --target packs --dry-run
  python scripts/fill_areeb_taxonomy.py --target cl --dry-run --fetch-orders
  python scripts/fill_areeb_taxonomy.py --target all
"""

from __future__ import annotations

import argparse
import csv
import json
import shutil
import sys
from collections import Counter, defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from shared import paths as wh
from shared.areeb_taxonomy import (
    AREEB_COLS,
    AreebCatalogs,
    AreebValues,
    apply_areeb,
    cell,
    classify_cl,
    load_catalogs,
)
from shared.cl_sku_match import build_label_index, match_keys, normalize_label

PLAIN_SHEET = "Sheet1"
PACKS_SHEET = "01-Database"
SS_STATUSES = ("shipped", "awaiting_shipment")


def backup_file(path: Path, dest_dir: Path) -> Path:
    dest_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = dest_dir / f"{path.stem}.bak_{stamp}{path.suffix}"
    shutil.copy2(path, dest)
    return dest


def header_index(headers: list[str]) -> dict[str, int]:
    return {h: i for i, h in enumerate(headers) if h}


def row_dict(headers: list[str], values: tuple) -> dict[str, object]:
    out: dict[str, object] = {}
    for i, h in enumerate(headers):
        if not h:
            continue
        out[h] = values[i] if i < len(values) else ""
    return out


def count_write(stats: dict[str, int], source: str, n_cells: int) -> None:
    stats["rows_touched"] += 1
    stats[f"src_{source or 'none'}"] += 1
    stats["cells_filled"] += n_cells


def fill_plain(cat: AreebCatalogs, *, dry_run: bool) -> dict[str, int]:
    path = wh.plain_database_path()
    wb = load_workbook(path, read_only=False, data_only=False)
    ws = wb[PLAIN_SHEET]
    headers = [cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers)
    for col in AREEB_COLS:
        if col not in idx:
            raise SystemExit(f"Plain Database missing column {col!r}")
    for col in ("SKU", "Product Code", "Brand", "Description"):
        if col not in idx:
            raise SystemExit(f"Plain Database missing column {col!r}")
    stats: dict[str, int] = defaultdict(int)
    leftover_desc: Counter[str] = Counter()
    samples: list[str] = []
    for r in range(2, ws.max_row + 1):
        sku = cell(ws.cell(r, idx["SKU"] + 1).value)
        current = {col: cell(ws.cell(r, idx[col] + 1).value) for col in AREEB_COLS}
        values = cat.classify_plain(
            sku,
            product_code=cell(ws.cell(r, idx["Product Code"] + 1).value),
            brand=cell(ws.cell(r, idx["Brand"] + 1).value),
            description=cell(ws.cell(r, idx["Description"] + 1).value),
        )
        patch = apply_areeb(current, values)
        stats["rows"] += 1
        if not patch:
            if not any(current.values()) and not values.any_filled():
                stats["blank_miss"] += 1
            else:
                stats["already_or_empty_incoming"] += 1
            continue
        count_write(stats, values.source, len(patch))
        if values.source == "plain_leftover":
            leftover_desc[cell(ws.cell(r, idx["Description"] + 1).value) or "(blank desc)"] += 1
        if len(samples) < 8:
            samples.append(f"SKU={sku} src={values.source} {patch}")
        if not dry_run:
            for col, val in patch.items():
                ws.cell(r, idx[col] + 1).value = val
    print("Plain Database:", path)
    _print_stats(stats, samples)
    if leftover_desc:
        print("  leftover Description top 12:")
        for k, n in leftover_desc.most_common(12):
            print(f"    {n:6,}  {k}")
    if dry_run:
        wb.close()
        return dict(stats)
    bak = backup_file(path, wh.plain_database_archive_dir())
    print(f"  backup {bak}")
    wb.save(path)
    wb.close()
    print("  wrote", path)
    return dict(stats)


def fill_packs(cat: AreebCatalogs, *, dry_run: bool) -> dict[str, int]:
    path = wh.packs_database_path()
    wb = load_workbook(path, read_only=False, data_only=False)
    ws = wb[PACKS_SHEET]
    headers = [cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers)
    for col in AREEB_COLS:
        if col not in idx:
            raise SystemExit(f"Packs Database missing column {col!r}")
    stats: dict[str, int] = defaultdict(int)
    samples: list[str] = []
    for r in range(2, ws.max_row + 1):
        item1 = cell(ws.cell(r, idx.get("Item 1 SKU", -1) + 1).value) if "Item 1 SKU" in idx else ""
        pcode = cell(ws.cell(r, idx.get("Product Code", -1) + 1).value) if "Product Code" in idx else ""
        child = (
            cell(ws.cell(r, idx.get("Channel Child SKU", -1) + 1).value)
            if "Channel Child SKU" in idx
            else ""
        )
        current = {col: cell(ws.cell(r, idx[col] + 1).value) for col in AREEB_COLS}
        values = cat.classify_packs(item1_sku=item1, product_code=pcode, channel_child_sku=child)
        patch = apply_areeb(current, values)
        stats["rows"] += 1
        if not patch:
            if not any(current.values()) and not values.any_filled():
                stats["blank_miss"] += 1
            else:
                stats["already_or_empty_incoming"] += 1
            continue
        count_write(stats, values.source, len(patch))
        if len(samples) < 8:
            samples.append(f"child={child} item1={item1} src={values.source} {patch}")
        if not dry_run:
            for col, val in patch.items():
                ws.cell(r, idx[col] + 1).value = val
    print("Packs Database:", path)
    _print_stats(stats, samples)
    if dry_run:
        wb.close()
        return dict(stats)
    bak = backup_file(path, wh.packs_database_archive_dir())
    print(f"  backup {bak}")
    wb.save(path)
    wb.close()
    print("  wrote", path)
    return dict(stats)


def _cl_rows(path: Path) -> tuple[list[str], list[dict[str, str]]]:
    with path.open(encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        headers = list(reader.fieldnames or [])
        rows = [{k: cell(v) for k, v in row.items()} for row in reader]
    return headers, rows


def fill_cl(*, dry_run: bool) -> tuple[dict[str, int], list[dict[str, str]]]:
    path = wh.cl_csv_path()
    headers, rows = _cl_rows(path)
    for col in AREEB_COLS:
        if col not in headers:
            raise SystemExit(f"CL missing column {col!r}")
    stats: dict[str, int] = defaultdict(int)
    samples: list[str] = []
    classified: list[tuple[dict[str, str], AreebValues]] = []
    for row in rows:
        values = classify_cl(row)
        patch = apply_areeb(row, values)
        stats["rows"] += 1
        classified.append((row, values))
        if not patch:
            if not any(cell(row.get(c)) for c in AREEB_COLS) and not values.any_filled():
                stats["blank_miss"] += 1
            else:
                stats["already_or_empty_incoming"] += 1
            continue
        count_write(stats, values.source, len(patch))
        if len(samples) < 8:
            samples.append(f"label={row.get('Custom Label')} src={values.source} {patch}")
        if not dry_run:
            row.update(patch)
    print("Custom Label:", path)
    _print_stats(stats, samples)
    leftover_rows = []
    for row, val in classified:
        if not val.category:
            leftover_rows.append(row)
    if dry_run:
        return dict(stats), leftover_rows
    bak = backup_file(path, wh.cl_backups_dir())
    print(f"  backup {bak}")
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)
    print("  wrote", path)
    return dict(stats), leftover_rows


def _print_stats(stats: dict[str, int], samples: list[str]) -> None:
    for k in sorted(stats):
        print(f"  {k}: {stats[k]:,}")
    for s in samples:
        print("  sample:", s[:220])


def ss_cache_path() -> Path:
    return wh.cl_app_dir() / "docs" / "areeb-taxonomy-ss-skus.json"


def extract_skus(orders: list[dict]) -> list[str]:
    skus: list[str] = []
    for order in orders:
        items = order.get("items") or []
        if not isinstance(items, list):
            continue
        for item in items:
            if not isinstance(item, dict):
                continue
            sku = cell(item.get("sku") or item.get("SKU"))
            if sku:
                skus.append(sku)
    return skus


def ss_page_cache_path() -> Path:
    return wh.cl_app_dir() / "docs" / "areeb-taxonomy-ss-pages.jsonl"


def _iter_page_cache(path: Path):
    """Yield parsed jsonl page records; skip blank/truncated lines."""
    if not path.is_file():
        return
    with path.open(encoding="utf-8-sig") as f:
        for line in f:
            line = line.strip()
            if not line:
                continue
            try:
                rec = json.loads(line)
            except json.JSONDecodeError:
                continue
            if not isinstance(rec, dict):
                continue
            yield rec


def _read_page_cache(path: Path) -> dict[str, set[int]]:
    done: dict[str, set[int]] = defaultdict(set)
    for rec in _iter_page_cache(path):
        done[str(rec.get("status") or "")].add(int(rec.get("page") or 0))
    return done


def fetch_year_skus(*, cache: Path, refresh: bool) -> list[str]:
    if cache.is_file() and not refresh:
        data = json.loads(cache.read_text(encoding="utf-8"))
        if data.get("complete"):
            skus = [cell(s) for s in data.get("skus") or [] if cell(s)]
            print(f"ShipStation SKU cache {cache} ({len(skus):,} SKUs)")
            return skus
    from shared.shipstation import ShipStationClient

    end = date.today()
    start = end - timedelta(days=365)
    start_s = start.isoformat()
    end_s = end.isoformat()
    page_cache = ss_page_cache_path()
    if refresh and page_cache.is_file():
        page_cache.unlink()
    done_pages = _read_page_cache(page_cache)
    print(f"Fetching ShipStation orders {start_s} .. {end_s} statuses={SS_STATUSES}")
    client = ShipStationClient(log=lambda m: print(" ", m, flush=True))
    all_skus: list[str] = []
    for rec in _iter_page_cache(page_cache):
        all_skus.extend(rec.get("skus") or [])
    if all_skus:
        print(f"  resumed {len(all_skus):,} SKUs from {page_cache.name}")

    page_cache.parent.mkdir(parents=True, exist_ok=True)
    counts: dict[str, int] = {}
    for status in SS_STATUSES:
        have = done_pages.get(status) or set()
        start_page = (max(have) + 1) if have else 1

        def on_page(page: int, total: int, batch: list, *, _status=status) -> None:
            skus = extract_skus(batch)
            all_skus.extend(skus)
            with page_cache.open("a", encoding="utf-8") as f:
                f.write(
                    json.dumps({"status": _status, "page": page, "pages": total, "skus": skus})
                    + "\n"
                )
            counts[_status] = counts.get(_status, 0) + len(batch)

        if have:
            print(f"  {status}: resume at page {start_page} (have {len(have)} pages)")
        client.list_orders(
            order_status=status,
            order_date_start=start_s,
            order_date_end=end_s,
            page_pause_s=2.5,
            start_page=start_page,
            collect=False,
            on_page=on_page,
        )
        print(f"  {status}: orders~{counts.get(status, 0):,}")
    unique = sorted(set(cell(s) for s in all_skus if cell(s)))
    cache.parent.mkdir(parents=True, exist_ok=True)
    cache.write_text(
        json.dumps(
            {
                "complete": True,
                "start": start_s,
                "end": end_s,
                "statuses": list(SS_STATUSES),
                "orders_by_status": counts,
                "sku_count": len(unique),
                "skus": unique,
            },
            indent=0,
        ),
        encoding="utf-8",
    )
    print(f"  cached {len(unique):,} unique SKUs -> {cache}")
    return unique


def analyze_orders_vs_cl(
    skus: list[str],
    cl_rows: list[dict[str, str]],
) -> None:
    labels = [r.get("Custom Label") or "" for r in cl_rows]
    index = build_label_index(labels, ids=range(len(cl_rows)))
    source_c = Counter()
    leftover_ga = Counter()
    leftover_cat = Counter()
    unmatched_sku = Counter()
    matched = 0
    for sku in skus:
        hit_id = None
        for key in match_keys(sku):
            hit_id = index.get(normalize_label(key))
            if hit_id is not None:
                break
        if hit_id is None:
            unmatched_sku[sku] += 1
            continue
        matched += 1
        row = cl_rows[int(hit_id)]
        values = classify_cl(row)
        source_c[values.source or "none"] += 1
        if not values.category:
            leftover_ga[cell(row.get("Gender Apparel")) or "(blank GA)"] += 1
            leftover_cat[cell(row.get("Category")) or "(blank Category)"] += 1
    print("\n=== Order SKUs vs CL (last year cache) ===")
    print(f"  unique order SKUs: {len(skus):,}")
    print(f"  matched a CL row: {matched:,}")
    print(f"  no CL row: {sum(unmatched_sku.values()):,} distinct={len(unmatched_sku):,}")
    print("  classify source on matched:")
    for k, n in source_c.most_common():
        print(f"    {k or 'none'}: {n:,}")
    print("  leftover still-blank Category (Areeb) Gender Apparel (top 15):")
    for k, n in leftover_ga.most_common(15):
        print(f"    {n:6,}  {k}")
    print("  leftover Category (top 10):")
    for k, n in leftover_cat.most_common(10):
        print(f"    {n:6,}  {k}")
    print("  unmatched order SKUs (top 15):")
    for k, n in unmatched_sku.most_common(15):
        print(f"    {n:6,}  {k}")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--target", choices=("plain", "packs", "cl", "all"), default="all")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--fetch-orders", action="store_true", help="Pull last 12 months of ShipStation SKUs")
    p.add_argument("--refresh-orders", action="store_true", help="Ignore SKU cache and fetch again")
    p.add_argument(
        "--orders-only",
        action="store_true",
        help="Skip catalog fills; fetch/analyze last-year SKUs only",
    )
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    dry = args.dry_run
    if dry:
        print("DRY RUN — no catalog writes\n", flush=True)
    targets = () if args.orders_only else (("plain", "packs", "cl") if args.target == "all" else (args.target,))
    cat: AreebCatalogs | None = None
    if "plain" in targets or "packs" in targets:
        print("Loading BTC + Uneek catalogs…", flush=True)
        cat = load_catalogs()
        print(
            f"  BTC UIDs={len(cat.btc_by_uid)//2:,} SPC={len(cat.btc_by_spc)//2:,} "
            f"Uneek short={len(cat.uneek_by_short)//2:,} code={len(cat.uneek_by_code)//2:,}",
            flush=True,
        )
    cl_rows_for_analysis: list[dict[str, str]] = []
    if "plain" in targets:
        assert cat is not None
        fill_plain(cat, dry_run=dry)
    if "packs" in targets:
        assert cat is not None
        fill_packs(cat, dry_run=dry)
    if "cl" in targets:
        _stats, leftover = fill_cl(dry_run=dry)
        _, cl_rows_for_analysis = _cl_rows(wh.cl_csv_path())
        ga = Counter(cell(r.get("Gender Apparel")) or "(blank)" for r in leftover)
        print("CL still-blank Category (Areeb) Gender Apparel top 12:")
        for k, n in ga.most_common(12):
            print(f"    {n:6,}  {k}")
    if args.fetch_orders or args.refresh_orders or args.orders_only:
        skus = fetch_year_skus(cache=ss_cache_path(), refresh=args.refresh_orders)
        if not cl_rows_for_analysis:
            _, cl_rows_for_analysis = _cl_rows(wh.cl_csv_path())
        analyze_orders_vs_cl(skus, cl_rows_for_analysis)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
