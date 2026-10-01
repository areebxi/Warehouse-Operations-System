"""Fill Areeb columns on Plain Database."""

from __future__ import annotations

from collections import Counter, defaultdict

from openpyxl import load_workbook

from fill_areeb_util import (
    PLAIN_SHEET,
    backup_file,
    count_write,
    header_index,
    print_stats,
)
from shared import paths as wh
from shared.areeb_taxonomy import AREEB_COLS, AreebCatalogs, apply_areeb, cell


def fill_plain(cat: AreebCatalogs, *, dry_run: bool) -> dict[str, int]:
    path = wh.plain_database_path()
    wb = load_workbook(path, read_only=False, data_only=False)
    ws = wb[PLAIN_SHEET]
    headers = [cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers)
    for col in AREEB_COLS:
        if col not in idx:
            raise SystemExit(f"Plain Database missing column {col!r}")
    for col in ("SKU", "Product Code", "Brand", "Description"):
        if col not in idx:
            raise SystemExit(f"Plain Database missing column {col!r}")
    stats: dict[str, int] = defaultdict(int)
    leftover_desc: Counter[str] = Counter()
    samples: list[str] = []
    for r in range(2, ws.max_row + 1):
        sku = cell(ws.cell(r, idx["SKU"] + 1).value)
        current = {col: cell(ws.cell(r, idx[col] + 1).value) for col in AREEB_COLS}
        values = cat.classify_plain(
            sku,
            product_code=cell(ws.cell(r, idx["Product Code"] + 1).value),
            brand=cell(ws.cell(r, idx["Brand"] + 1).value),
            description=cell(ws.cell(r, idx["Description"] + 1).value),
        )
        patch = apply_areeb(current, values)
        stats["rows"] += 1
        if not patch:
            if not any(current.values()) and not values.any_filled():
                stats["blank_miss"] += 1
            else:
                stats["already_or_empty_incoming"] += 1
            continue
        count_write(stats, values.source, len(patch))
        if values.source == "plain_leftover":
            leftover_desc[cell(ws.cell(r, idx["Description"] + 1).value) or "(blank desc)"] += 1
        if len(samples) < 8:
            samples.append(f"SKU={sku} src={values.source} {patch}")
        if not dry_run:
            for col, val in patch.items():
                ws.cell(r, idx[col] + 1).value = val
    print("Plain Database:", path)
    print_stats(stats, samples)
    if leftover_desc:
        print("  leftover Description top 12:")
        for k, n in leftover_desc.most_common(12):
            print(f"    {n:6,}  {k}")
    if dry_run:
        wb.close()
        return dict(stats)
    bak = backup_file(path, wh.plain_database_archive_dir())
    print(f"  backup {bak}")
    wb.save(path)
    wb.close()
    print("  wrote", path)
    return dict(stats)
