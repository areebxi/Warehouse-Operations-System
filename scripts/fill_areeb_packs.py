"""Fill Areeb columns on Packs Database."""

from __future__ import annotations

from collections import defaultdict

from openpyxl import load_workbook

from fill_areeb_util import (
    PACKS_SHEET,
    backup_file,
    count_write,
    header_index,
    print_stats,
)
from shared import paths as wh
from shared.areeb_taxonomy import AREEB_COLS, AreebCatalogs, apply_areeb, cell


def fill_packs(cat: AreebCatalogs, *, dry_run: bool) -> dict[str, int]:
    path = wh.packs_database_path()
    wb = load_workbook(path, read_only=False, data_only=False)
    ws = wb[PACKS_SHEET]
    headers = [cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    idx = header_index(headers)
    for col in AREEB_COLS:
        if col not in idx:
            raise SystemExit(f"Packs Database missing column {col!r}")
    stats: dict[str, int] = defaultdict(int)
    samples: list[str] = []
    for r in range(2, ws.max_row + 1):
        item1 = cell(ws.cell(r, idx.get("Item 1 SKU", -1) + 1).value) if "Item 1 SKU" in idx else ""
        pcode = cell(ws.cell(r, idx.get("Product Code", -1) + 1).value) if "Product Code" in idx else ""
        child = (
            cell(ws.cell(r, idx.get("Channel Child SKU", -1) + 1).value)
            if "Channel Child SKU" in idx
            else ""
        )
        current = {col: cell(ws.cell(r, idx[col] + 1).value) for col in AREEB_COLS}
        values = cat.classify_packs(item1_sku=item1, product_code=pcode, channel_child_sku=child)
        patch = apply_areeb(current, values)
        stats["rows"] += 1
        if not patch:
            if not any(current.values()) and not values.any_filled():
                stats["blank_miss"] += 1
            else:
                stats["already_or_empty_incoming"] += 1
            continue
        count_write(stats, values.source, len(patch))
        if len(samples) < 8:
            samples.append(f"child={child} item1={item1} src={values.source} {patch}")
        if not dry_run:
            for col, val in patch.items():
                ws.cell(r, idx[col] + 1).value = val
    print("Packs Database:", path)
    print_stats(stats, samples)
    if dry_run:
        wb.close()
        return dict(stats)
    bak = backup_file(path, wh.packs_database_archive_dir())
    print(f"  backup {bak}")
    wb.save(path)
    wb.close()
    print("  wrote", path)
    return dict(stats)
