"""
Sync Database Transfer xlsx sheets from live CL + Size References CSVs.

  Database Transfer/Workbook.xlsx          → sheet "CL Database"
  Database Transfer/Configuration Workbook.xlsx → sheet "Size References"

Other Workbook sheets and Override Print Size are preserved.
Run after a supervisor fill of CL and/or Size References.

  python scripts/sync_database_transfer.py
  python scripts/sync_database_transfer.py --dry-run
"""
from __future__ import annotations

import argparse
import shutil
import sys
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook

_SCRIPT_DIR = Path(__file__).resolve().parent
_APP_ROOT = _SCRIPT_DIR.parent
_WAREHOUSE_ROOT = _APP_ROOT.parent
sys.path.insert(0, str(_WAREHOUSE_ROOT))

from shared.paths import (  # noqa: E402
    cl_csv_path,
    database_transfer_config_workbook_path,
    database_transfer_backups_dir,
    database_transfer_workbook_path,
    size_references_csv_path,
)

# Transfer Workbook uses older NocoDB-facing aliases for three columns.
_CL_RENAME = {
    "Colour": "Colour Name",
    "Apparel Image": "Picture Name",
    "Print Positions": "Position",
}


def _sheet_to_df(ws) -> pd.DataFrame:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame()
    header = ["" if c is None else str(c) for c in rows[0]]
    data = []
    for r in rows[1:]:
        cells = [("" if c is None else c) for c in r]
        if len(cells) < len(header):
            cells = list(cells) + [""] * (len(header) - len(cells))
        data.append(cells[: len(header)])
    return pd.DataFrame(data, columns=header)


def _write_df_sheet(wb: Workbook, title: str, df: pd.DataFrame) -> None:
    ws = wb.create_sheet(title)
    ws.append(list(df.columns))
    for row in df.itertuples(index=False, name=None):
        ws.append([("" if v is None else v) for v in row])


def _backup(path: Path) -> Path:
    bak_dir = database_transfer_backups_dir()
    bak_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    dest = bak_dir / f"{path.stem}_preSync_{stamp}{path.suffix}"
    shutil.copy2(path, dest)
    return dest


def sync(*, dry_run: bool = False) -> dict:
    cl_path = cl_csv_path()
    sr_path = size_references_csv_path()
    wb_path = database_transfer_workbook_path()
    cfg_path = database_transfer_config_workbook_path()

    cl = pd.read_csv(cl_path, dtype=str, encoding="utf-8-sig", keep_default_na=False)
    sr = pd.read_csv(sr_path, dtype=str, encoding="utf-8-sig", keep_default_na=False)
    cl_out = cl.rename(columns=_CL_RENAME)
    cl_out = cl_out[[_CL_RENAME.get(c, c) for c in cl.columns]]

    summary = {
        "cl_rows": len(cl_out),
        "sr_rows": len(sr),
        "cl_cols": list(cl_out.columns),
        "sr_cols": list(sr.columns),
        "dry_run": dry_run,
    }
    if dry_run:
        return summary

    bak_wb = _backup(wb_path)
    bak_cfg = _backup(cfg_path)
    summary["backups"] = [str(bak_wb), str(bak_cfg)]

    old_cfg = load_workbook(cfg_path, read_only=True, data_only=True)
    override = _sheet_to_df(old_cfg["Override Print Size"])
    old_cfg.close()

    cfg_wb = Workbook(write_only=True)
    _write_df_sheet(cfg_wb, "Size References", sr)
    _write_df_sheet(cfg_wb, "Override Print Size", override)
    tmp_cfg = cfg_path.with_suffix(".xlsx.tmp")
    cfg_wb.save(tmp_cfg)
    tmp_cfg.replace(cfg_path)

    old_wb = load_workbook(wb_path, read_only=True, data_only=True)
    other_order = [s for s in old_wb.sheetnames if s != "CL Database"]
    others = {name: _sheet_to_df(old_wb[name]) for name in other_order}
    old_wb.close()

    wb = Workbook(write_only=True)
    _write_df_sheet(wb, "CL Database", cl_out)
    for name in other_order:
        _write_df_sheet(wb, name, others[name])
    tmp_wb = wb_path.with_suffix(".xlsx.tmp")
    wb.save(tmp_wb)
    tmp_wb.replace(wb_path)

    summary["override_rows"] = len(override)
    summary["preserved_sheets"] = other_order
    return summary


def main() -> int:
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument("--dry-run", action="store_true")
    args = p.parse_args()
    summary = sync(dry_run=args.dry_run)
    mode = "dry-run" if summary["dry_run"] else "wrote"
    print(f"{mode}: CL Database {summary['cl_rows']} rows; Size References {summary['sr_rows']} rows")
    if summary.get("backups"):
        for b in summary["backups"]:
            print(f"  backup: {b}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
