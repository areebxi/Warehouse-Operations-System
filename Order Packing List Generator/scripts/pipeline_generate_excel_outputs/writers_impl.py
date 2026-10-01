from __future__ import annotations
from datetime import date
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from .helpers import _extended_process_and_item_number, _gender_colour_size_combo_hyphenated, _item_number_from_extended, _normalize, _order_number_base, _process_number_for_excel_from_row, _process_plus_additional, _remap_dtf_item_sku, _split_item_sku_by_lg

def _write_dtf_des(
    df: pd.DataFrame,
    path: Path,
    use_fixed_process_number: bool = False,
    use_fixed_numeric_process: bool = False,
    dtf_sku_map: dict[str, str] | None = None,
) -> None:
    headers = [
        "Order - Number", "Item - Qty", "Item - SKU", "Item - Name", "Ship To - Name",
        "Notes - From Buyer", "Ship To - Postal Code", "Source", "Process Num", "Genre",
        "Order Type", "Orders Type Abbrevation", "Condition", "Customise", "", "Item Num",
    ]
    if df.empty:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        for c, h in enumerate(headers, start=1):
            ws.cell(row=1, column=c, value=h)
        wb.save(path)
        return

    rows = []
    for _, row in df.iterrows():
        process_and_item = row.get("Process and Item Number", "")
        extended = _extended_process_and_item_number(process_and_item)
        process_excel = _process_number_for_excel_from_row(process_and_item)
        process_display = f"Process {process_excel}" if process_excel else ""
        item_num_str = _item_number_from_extended(extended)
        item_display = f"Item {item_num_str}" if item_num_str else "Item 1"
        item_sku = _normalize(row.get("Item SKU", ""))
        customise_val = _normalize(row.get("Customise", ""))
        is_customised = customise_val.lower() == "yes"
        segments = [item_sku] if is_customised else _split_item_sku_by_lg(item_sku)
        m = dtf_sku_map or {}
        segments = [_remap_dtf_item_sku(seg, m) for seg in segments]
        qty = max(1, int(row.get("Item Quantity", 1) or 1))
        notes_from_buyer = (
            _normalize(row.get("Notes From Buyer", ""))
            or _normalize(row.get("Notes - From Buyer", ""))
        )
        base_row = [
            _order_number_base(row),
            1,
            item_sku,
            _normalize(row.get("Item Name", "")),
            _normalize(row.get("Recipient Name", "")),
            notes_from_buyer, "", "", process_display, _normalize(row.get("Gender Apparel", "")),
            "", "", "",
        ]
        for seg in segments:
            for _ in range(qty):
                row_data = base_row.copy()
                row_data[2] = seg
                row_data.extend([customise_val, "", item_display])
                rows.append(row_data)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    for r, row_data in enumerate(rows, start=2):
        for c, val in enumerate(row_data, start=1):
            ws.cell(row=r, column=c, value=val)
    wb.save(path)
