from datetime import date
from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from .helpers import _extended_process_and_item_number, _gender_colour_size_combo_hyphenated, _item_number_from_extended, _normalize, _order_number_base, _process_number_for_excel_from_row, _process_plus_additional, _remap_dtf_item_sku, _split_item_sku_by_lg
from .writers_picking import (
    _write_picking,
)
def _write_orders_details(df: pd.DataFrame, process_base: str, path: Path) -> None:
    if df.empty:
        wb = Workbook()
        ws = wb.active
        ws.title = "Sheet1"
        wb.save(path)
        return

    order_counts = df["Order Number (Base)"].fillna("").astype(str).value_counts()
    df = df.copy()
    df["_extended"] = df["Process and Item Number"].map(_extended_process_and_item_number)
    df["_block"] = df["_extended"].map(_process_plus_additional)
    block_recipient_count = df.groupby("_block")["Recipient Name"].nunique().to_dict()
    df["_qty"] = df["Item Quantity"].fillna(1).astype(int)
    block_quantity = df.groupby("_block")["_qty"].sum().to_dict()

    def _row_is_merge(r: pd.Series) -> bool:
        order_num = _order_number_base(r)
        is_repeated_order = order_counts.get(order_num, 0) >= 2
        has_multi_qty = r["_qty"] > 1
        return is_repeated_order or has_multi_qty

    df["_is_merge_row"] = df.apply(_row_is_merge, axis=1)

    seen = set()
    blocks_ordered = []
    for b in df["_block"]:
        if b not in seen:
            seen.add(b)
            blocks_ordered.append(b)

    headers = [
        "Condition", "Gender Apparel-Colour Name-Size", "Merge", "Single", "Single",
        "Porcess Number", "Gender Apparel-Colour Name-Size", "H", "Type", "J", "K",
        "Condition", "Porcess Number", "N", "O", "Gender Apparel-Colour Name-Size", "Q",
        "Porcess Number", "Merge", "Single",
    ]
    data = []
    for block in blocks_ordered:
        grp = df[df["_block"] == block]
        first = grp.iloc[0]
        is_merge = bool(grp["_is_merge_row"].any())
        merge_val = block_recipient_count.get(block, 0)
        single_val = int(block_quantity.get(block, 0))

        cond = "Condition 1 Merge" if is_merge else "Condition 4"
        combo = "Merge Orders" if is_merge else _gender_colour_size_combo_hyphenated(first)
        customise = _normalize(first.get("Customise", ""))
        type_base = "Personalised" if customise.lower() == "yes" else "Normal"
        i_val = f"{type_base}-Merge" if is_merge else type_base

        process_excel = _process_number_for_excel_from_row(first.get("Process and Item Number", ""))
        data.append([
            cond, combo, merge_val, single_val, single_val,
            process_excel, combo, "", i_val, "", "", cond, process_excel, "", "", combo, "",
            process_excel, merge_val, single_val,
        ])

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for c, h in enumerate(headers, start=1):
        ws.cell(row=1, column=c, value=h)
    for r, row_data in enumerate(data, start=2):
        for c, val in enumerate(row_data, start=1):
            ws.cell(row=r, column=c, value=val)
    wb.save(path)
