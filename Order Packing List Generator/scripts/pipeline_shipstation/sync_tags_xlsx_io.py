"""Excel IO helpers for ShipStation Tags.xlsx sync."""
from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

SHEET_NAME = "Tags"
COL_SR = 1
COL_NAME = 2
COL_ID = 3
HEADERS = [
    "Sr. No.",
    "Tag Name",
    "Tag ID",
    "Process No - 1st Shift",
    "Process No - 2nd Shift",
    "Process No - 3rd Shift",
    "Process No - 4th Shift",
    "Process No - 5th Shift",
]


def _name_key(name: str) -> str:
    return str(name or "").strip().casefold()


def _as_int(value: Any) -> int | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    try:
        return int(float(text)) if "." in text else int(text)
    except (TypeError, ValueError):
        return None


def _ensure_headers(ws: Worksheet) -> None:
    for col, header in enumerate(HEADERS, start=1):
        current = ws.cell(1, col).value
        if current is None or str(current).strip() == "":
            ws.cell(1, col).value = header


def _read_excel_rows(ws: Worksheet) -> dict[str, tuple[int, str, int | None]]:
    """Map casefolded tag name -> (row, display name, tag id)."""
    by_name: dict[str, tuple[int, str, int | None]] = {}
    for row in range(2, ws.max_row + 1):
        name_raw = ws.cell(row, COL_NAME).value
        if name_raw is None or str(name_raw).strip() == "":
            continue
        name = str(name_raw).strip()
        key = _name_key(name)
        by_name[key] = (row, name, _as_int(ws.cell(row, COL_ID).value))
    return by_name


def _next_sr_no(ws: Worksheet) -> int:
    last = 0
    for row in range(2, ws.max_row + 1):
        n = _as_int(ws.cell(row, COL_SR).value)
        if n is not None and n > last:
            last = n
    return last + 1


def _create_empty_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    for col, header in enumerate(HEADERS, start=1):
        ws.cell(1, col).value = header
    wb.save(path)
