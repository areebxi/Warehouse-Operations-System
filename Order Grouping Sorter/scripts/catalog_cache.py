"""Disk cache for Plain / Packs catalog indexes (mtime + size invalidate)."""

from __future__ import annotations

import csv
import pickle
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from shared import paths as wh
from shared.areeb_taxonomy import cell

CACHE_VERSION = 1


def _fold(value: object) -> str:
    return cell(value).casefold()


def _source_meta(path: Path, sheet: str, key_col: str) -> dict[str, Any]:
    st = path.stat()
    return {
        "path": str(path.resolve()),
        "mtime_ns": st.st_mtime_ns,
        "size": st.st_size,
        "sheet": sheet,
        "key_col": key_col,
    }


def _index_rows(rows: list[dict[str, str]], key_col: str) -> dict[str, dict[str, str]]:
    idx: dict[str, dict[str, str]] = {}
    for row in rows:
        key = _fold(row.get(key_col))
        if key and key not in idx:
            idx[key] = row
    return idx


def _read_csv_index(path: Path, key_col: str) -> dict[str, dict[str, str]]:
    last: Exception | None = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with path.open(encoding=enc, newline="") as f:
                rows = [{k: cell(v) for k, v in row.items() if k} for row in csv.DictReader(f)]
            return _index_rows(rows, key_col)
        except UnicodeDecodeError as exc:
            last = exc
    raise last or RuntimeError(f"Could not decode {path}")


def read_xlsx_index(path: Path, sheet: str, key_col: str) -> dict[str, dict[str, str]]:
    if path.suffix.lower() == ".csv":
        return _read_csv_index(path, key_col)
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        it = ws.iter_rows(values_only=True)
        headers = [cell(h) for h in next(it)]
        rows: list[dict[str, str]] = []
        for raw in it:
            row = {
                headers[i]: cell(raw[i]) if i < len(raw) else ""
                for i in range(len(headers))
                if headers[i]
            }
            rows.append(row)
    finally:
        wb.close()
    return _index_rows(rows, key_col)


def load_xlsx_index(
    path: Path,
    sheet: str,
    key_col: str,
    *,
    cache_stem: str,
) -> dict[str, dict[str, str]]:
    """Load keyed index; reuse pickle when source file meta matches."""
    cache_path = wh.sorter_catalog_cache_path(cache_stem)
    meta = _source_meta(path, sheet, key_col)
    if cache_path.is_file():
        try:
            payload = pickle.loads(cache_path.read_bytes())
            if (
                isinstance(payload, dict)
                and payload.get("v") == CACHE_VERSION
                and payload.get("meta") == meta
                and isinstance(payload.get("index"), dict)
            ):
                return payload["index"]
        except Exception:
            pass  # ponytail: corrupt cache → rebuild; ceiling = rare disk glitch
    index = read_xlsx_index(path, sheet, key_col)
    cache_path.parent.mkdir(parents=True, exist_ok=True)
    cache_path.write_bytes(
        pickle.dumps({"v": CACHE_VERSION, "meta": meta, "index": index}, protocol=pickle.HIGHEST_PROTOCOL)
    )
    return index
