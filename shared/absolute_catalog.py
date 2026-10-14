"""Absolute Product Data index, same role as BTC / Uneek catalogs.

Match is the sheet Stock Code (``AA22-36-L``), not the style alone.
A style that only appears on this price list does not steal a BTC UID.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path

from shared import paths as wh
from shared.areeb_taxonomy.cl_classify import cl_standard
from shared.areeb_taxonomy.values import cell

ABSOLUTE_APPARELS = "Absolute Apparels"

ON_DEMAND = "Supplier On Demand"


def absolute_source_path(from_path: object | None = None) -> Path:
    """CSV price list when the supervisor has added one; else the xlsx path."""
    xlsx = wh.absolute_product_data_path(from_path)
    csv_path = xlsx.with_suffix(".csv")
    if csv_path.is_file():
        return csv_path
    return xlsx


def load_absolute_rows(path: Path | None = None) -> list[dict[str, str]]:
    src = path or absolute_source_path()
    if src.suffix.lower() == ".csv":
        return _read_csv(src)
    return _read_xlsx(src)


def index_stock_codes(rows: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    """Stock Code → grouping row. Original case and casefold both point at it."""
    idx: dict[str, dict[str, str]] = {}
    for raw in rows:
        code = cell(raw.get("Stock Code"))
        if not code or code in idx:
            continue
        row = grouping_row(raw)
        idx[code] = row
        idx[code.casefold()] = row
    return idx


def _classify_text(desc: str) -> str:
    """Absolute descriptions say Hood / Sweat. The warehouse classifier says Hoodie / Sweatshirt."""
    text = desc
    if re.search(r"\bhood\b", text, re.I) and "hoodie" not in text.casefold() and "hooded" not in text.casefold():
        text = re.sub(r"\bhood\b", "Hoodie", text, count=1, flags=re.I)
    if re.search(r"\bsweat\b", text, re.I) and "sweatshirt" not in text.casefold():
        text = re.sub(r"\bsweat\b", "Sweatshirt", text, count=1, flags=re.I)
    return text


def grouping_row(raw: dict[str, str]) -> dict[str, str]:
    """Plain-catalog shape so the sorter can place an Absolute stock code."""
    desc = cell(raw.get("Description"))
    areeb = cl_standard({"Gender Apparel": _classify_text(desc)}).as_dict()
    if areeb.get("Category (Areeb)") and not areeb.get("Product Style (Areeb)"):
        areeb["Product Style (Areeb)"] = "Standard"
    row = {
        "SKU": cell(raw.get("Stock Code")),
        "Product Code": cell(raw.get("Style")),
        "Brand": cell(raw.get("Supplier")),
        "Colour": cell(raw.get("Colour Name")),
        "Size": cell(raw.get("Size Name")),
        "Description": desc,
        "Supply Method": ON_DEMAND,
        "Supplier Name": ABSOLUTE_APPARELS,
        "Printing Type": "",
    }
    row.update(areeb)
    return row


def _read_csv(path: Path) -> list[dict[str, str]]:
    last: Exception | None = None
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            with path.open(encoding=enc, newline="") as f:
                return [{k: cell(v) for k, v in row.items() if k} for row in csv.DictReader(f)]
        except UnicodeDecodeError as exc:
            last = exc
    raise last or RuntimeError(f"Could not decode {path}")


def _read_xlsx(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        headers = [cell(h) for h in next(it)]
        out: list[dict[str, str]] = []
        for raw in it:
            row = {
                headers[i]: cell(raw[i] if i < len(raw) else "")
                for i in range(len(headers))
                if headers[i]
            }
            if any(row.values()):
                out.append(row)
        return out
    finally:
        wb.close()
