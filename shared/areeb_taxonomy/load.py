"""Load BTC/Uneek catalogs and apply Areeb patches to rows."""
from __future__ import annotations

from pathlib import Path
from typing import Any, Mapping

from shared.areeb_taxonomy.catalogs import AreebCatalogs
from shared.areeb_taxonomy.consts import AREEB_COLS, SOURCE_CL_STANDARD
from shared.areeb_taxonomy.values import AreebValues, cell

def _norm_row(row: Mapping[str, Any]) -> dict[str, str]:
    return {str(k): cell(v) for k, v in row.items()}


def load_btc_into(cat: AreebCatalogs, rows: list[Mapping[str, Any]]) -> None:
    for raw in rows:
        row = _norm_row(raw)
        uid = cell(row.get("UID") or row.get("Sku") or row.get("SKU"))
        if not uid or uid.startswith("["):
            continue
        if uid not in cat.btc_by_uid:
            cat.btc_by_uid[uid] = row
            cat.btc_by_uid[uid.casefold()] = row
        spc = cell(row.get("SPC"))
        if spc and spc not in cat.btc_by_spc:
            cat.btc_by_spc[spc] = row
            cat.btc_by_spc[spc.casefold()] = row


def load_uneek_into(cat: AreebCatalogs, rows: list[Mapping[str, Any]]) -> None:
    for raw in rows:
        row = _norm_row(raw)
        short = cell(row.get("Short Code"))
        if short and short not in cat.uneek_by_short:
            cat.uneek_by_short[short] = row
            cat.uneek_by_short[short.casefold()] = row
        code = cell(row.get("Product Code"))
        if code and code not in cat.uneek_by_code:
            cat.uneek_by_code[code] = row
            cat.uneek_by_code[code.casefold()] = row


def load_btc_csv(path: Path) -> list[dict[str, str]]:
    import csv

    last_err: Exception | None = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with path.open(encoding=enc, newline="") as f:
                return [{k: cell(v) for k, v in row.items()} for row in csv.DictReader(f)]
        except UnicodeDecodeError as exc:
            last_err = exc
            continue
    raise last_err or RuntimeError(f"Could not decode {path}")


def load_uneek_xlsx(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        headers = [cell(h) for h in next(rows_iter)]
        out: list[dict[str, str]] = []
        for raw in rows_iter:
            row = {headers[i]: cell(raw[i] if i < len(raw) else "") for i in range(len(headers)) if headers[i]}
            if any(row.values()):
                out.append(row)
        return out
    finally:
        wb.close()


def load_catalogs(
    *,
    btc_path: Path | None = None,
    uneek_path: Path | None = None,
) -> AreebCatalogs:
    from shared import paths as wh

    cat = AreebCatalogs()
    load_btc_into(cat, load_btc_csv(btc_path or wh.btc_product_data_path()))
    load_uneek_into(cat, load_uneek_xlsx(uneek_path or wh.uneek_product_data_path()))
    return cat


def apply_blank_only(current: Mapping[str, Any], values: AreebValues) -> dict[str, str]:
    """Return Areeb cells to write: only where current is blank and incoming is non-blank."""
    out: dict[str, str] = {}
    incoming = values.as_dict()
    for col in AREEB_COLS:
        if cell(current.get(col)):
            continue
        val = incoming[col]
        if val:
            out[col] = val
    return out


def apply_areeb(current: Mapping[str, Any], values: AreebValues) -> dict[str, str]:
    """Blank-only for supplier joins. CL standard overwrites all four Areeb cells.

    Hashim #038: an off-list warehouse style/type/category is written blank
    (do not keep an invented cell).
    """
    if values.source != SOURCE_CL_STANDARD:
        return apply_blank_only(current, values)
    out: dict[str, str] = {}
    for col, val in values.as_dict().items():
        new = cell(val)
        if cell(current.get(col)) != new:
            out[col] = new
    return out
