"""Four Areeb 30-chain columns: maps + classify.

Plain / Packs: blank-only copy from BTC / Uneek.
Plain leftover (no UID / Short Code / SPC): Brand + Description; Category/Type from Description in BTC language.
Custom Label: warehouse `cl_standard` from Gender Apparel only. Never BTC, Uneek, Brand, or PE cells.
CL fill snaps to Hashim #038 Title Case pick-lists (type has no gender; style is a product name).
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Mapping

import shared.taxonomy_catalog as tax
from shared.taxonomy_picklist import snap_warehouse

AREEB_COLS = (
    "Category (Areeb)",
    "Product Type (Areeb)",
    "Product Style (Areeb)",
    "Department (Areeb)",
)

BTC_SRC = {
    "Category (Areeb)": "Department",
    "Product Type (Areeb)": "Sub Department",
    "Product Style (Areeb)": "Brand",
    "Department (Areeb)": "Description",
}

UNEEK_SRC = {
    "Category (Areeb)": "Category",
    "Product Type (Areeb)": "Product Name",
    "Product Style (Areeb)": "Full Description",
    "Department (Areeb)": "Gender",
}

GENDER_MENS = "Mens"
GENDER_WOMENS = "Womens"
GENDER_KIDS = "Kids"
GENDER_UNISEX = "Unisex"
GENDER_GENERAL = "General"

# Exact Gender Apparel → warehouse 4-tuple. Overrides the pattern engine.
# Tuple: Category, Product Type, Product Style, Department (gender only).
CL_STANDARD_RULES: dict[str, tuple[str, str, str, str]] = {
    "Mens-T-Shirt": (tax.CAT_TEE, tax.TYPE_SS_TEE, tax.DEFAULT_STYLE, GENDER_MENS),
    "Womens-T-Shirt": (tax.CAT_TEE, tax.TYPE_SS_TEE, tax.DEFAULT_STYLE, GENDER_WOMENS),
    "Kids-T-Shirt": (tax.CAT_TEE, tax.TYPE_SS_TEE, tax.DEFAULT_STYLE, GENDER_KIDS),
    "GILDAN Heavy Cotton Adult T-Shirt": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Heavy Cotton", GENDER_MENS),
    "GILDAN Softstyle Adult T-Shirt": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Softstyle", GENDER_MENS),
    "GILDAN Softstyle Ladies T Shirt": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Softstyle", GENDER_WOMENS),
    "GILDAN Softstyle Ladies Tank Top": (tax.CAT_TEE, tax.TYPE_TANK, "Softstyle", GENDER_WOMENS),
    "FOTL Kids Valueweight T": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Valueweight", GENDER_KIDS),
    "5000": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Heavy Cotton", GENDER_MENS),
    "G5000": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Heavy Cotton", GENDER_MENS),
    "2400": (tax.CAT_TEE, tax.TYPE_LS_TEE, "Ultra Cotton", GENDER_MENS),
    "61026": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Valueweight Baseball", GENDER_MENS),
    "61168-Ringer": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Ringer", GENDER_MENS),
    "TD02B": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Kids Tie Dye", GENDER_KIDS),
    "CA3001T": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Toddler Jersey", GENDER_KIDS),
    "JC003": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Contrast Cool", GENDER_MENS),
    "JC03J": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Contrast Cool", GENDER_KIDS),
    "Acid Wash Vintage Rust": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Acid Wash Vintage Rust", GENDER_MENS),
    "Acid Wash Optic Wash": (tax.CAT_TEE, tax.TYPE_SS_TEE, "Acid Wash Optic Wash", GENDER_MENS),
    "Mens-Hoodie": (tax.CAT_SWEAT, tax.TYPE_HOODIE, tax.DEFAULT_STYLE, GENDER_MENS),
    "Mens-Sweatshirt": (tax.CAT_SWEAT, tax.TYPE_SWEAT, tax.DEFAULT_STYLE, GENDER_MENS),
    "Womens-Hoodie": (tax.CAT_SWEAT, tax.TYPE_HOODIE, tax.DEFAULT_STYLE, GENDER_WOMENS),
    "Womens-Sweatshirt": (tax.CAT_SWEAT, tax.TYPE_SWEAT, tax.DEFAULT_STYLE, GENDER_WOMENS),
    "Kids-Hoodie": (tax.CAT_SWEAT, tax.TYPE_HOODIE, tax.DEFAULT_STYLE, GENDER_KIDS),
    "Kids-Sweatshirt": (tax.CAT_SWEAT, tax.TYPE_SWEAT, tax.DEFAULT_STYLE, GENDER_KIDS),
    "GILDAN Heavy Blend Adult Hooded Sweatshirt": (tax.CAT_SWEAT, tax.TYPE_HOODIE, "Heavy Blend", GENDER_MENS),
    "GILDAN Softstyle Midw Fleece Youth Hoodie": (
        tax.CAT_SWEAT,
        tax.TYPE_HOODIE,
        "Softstyle Midweight Fleece",
        GENDER_KIDS,
    ),
    "GILDAN Softstyle Midweight Fleece Youth Hoodie": (
        tax.CAT_SWEAT,
        tax.TYPE_HOODIE,
        "Softstyle Midweight Fleece",
        GENDER_KIDS,
    ),
    "JH001-Hoodie": (tax.CAT_SWEAT, tax.TYPE_HOODIE, "College", GENDER_MENS),
    "JH01J-Hoodie": (tax.CAT_SWEAT, tax.TYPE_HOODIE, "College", GENDER_KIDS),
    "C2200-Hoodie": (tax.CAT_SWEAT, tax.TYPE_HOODIE, "Ringspun Blended", GENDER_MENS),
    "C2400": (tax.CAT_SWEAT, tax.TYPE_SWEAT, "Ringspun Blended", GENDER_MENS),
    "SF500B-Hoodie": (tax.CAT_SWEAT, tax.TYPE_HOODIE, "Softstyle Midweight Fleece", GENDER_KIDS),
    "Kids-T-Shirt-Hoodie": (tax.CAT_SETS, tax.TYPE_SET, "Kids Set", GENDER_KIDS),
    "85800L-Polo-T-Shirt": (tax.CAT_POLO, tax.TYPE_SS_POLO, "Premium Cotton Polo", GENDER_WOMENS),
    "64800-Polo T-Shirt": (tax.CAT_POLO, tax.TYPE_SS_POLO, "Softstyle Double Pique", GENDER_MENS),
    "UCC003-Mens-Polo": (tax.CAT_POLO, tax.TYPE_SS_POLO, "Everyday Polo", GENDER_MENS),
    "Uneek Hi-Viz Polo Shirt": (tax.CAT_POLO, tax.TYPE_HV_POLO, "High Visibility", GENDER_UNISEX),
    "64200": (tax.CAT_TEE, tax.TYPE_TANK, "Softstyle", GENDER_MENS),
    "64200L": (tax.CAT_TEE, tax.TYPE_TANK, "Softstyle", GENDER_WOMENS),
    "C800T-BS": (tax.CAT_BABY, tax.TYPE_BODY, "Baby Body Suit", GENDER_KIDS),
    "C8030T-BS": (tax.CAT_BABY, tax.TYPE_ROMPER, "Baby Romper", GENDER_KIDS),
    "BZ02-Toddler-T-Shirt": (tax.CAT_BABY, tax.TYPE_BABY_TEE, "Baby T-Shirt", GENDER_KIDS),
    "BZ10-Body Suit": (tax.CAT_BABY, tax.TYPE_BODY, "Baby Bodysuit", GENDER_KIDS),
    "Kids-Toddler 61033": (tax.CAT_BABY, tax.TYPE_BABY_TEE, "Valueweight", GENDER_KIDS),
    "China Bag": (tax.CAT_BAGS, tax.TYPE_TOTE, "China Bag", GENDER_GENERAL),
    "Cotton-Shopper": (tax.CAT_BAGS, tax.TYPE_TOTE, "Cotton Shopper", GENDER_GENERAL),
    "BagBase Boutique Wristlet Keyring": (
        tax.CAT_BAGS,
        tax.TYPE_BAG_ACC,
        "Boutique Wristlet Keyring",
        GENDER_GENERAL,
    ),
    "PC-QD442": (tax.CAT_BAGS, tax.TYPE_PENCIL, "Pencil Case", GENDER_GENERAL),
    "WB-QD440": (tax.CAT_BAGS, tax.TYPE_BAG_ACC, "Water Bottle Holder", GENDER_GENERAL),
    "W696": (tax.CAT_BAGS, tax.TYPE_TOTE, "Oversized Canvas Tote", GENDER_GENERAL),
    "W265": (tax.CAT_BAGS, tax.TYPE_TOTE, "Organic Premium Maxi Tote", GENDER_GENERAL),
    "BG745": (tax.CAT_BAGS, tax.TYPE_BAG_ACC, "Boutique Circular Key Clip", GENDER_GENERAL),
    "Yoko Hi-Vis Class 2 Waistcoat": (tax.CAT_SAFETY, tax.TYPE_HV_VEST, "Class 2", GENDER_UNISEX),
    "Hi-Vis-HVW801": (tax.CAT_SAFETY, tax.TYPE_HV_VEST, "Executive Vest", GENDER_UNISEX),
    "Beechfield Original Patch Beanie": (tax.CAT_HEAD, tax.TYPE_BEANIE, "Original Patch", GENDER_GENERAL),
    "Beechfield Snowstar Patch Beanie": (tax.CAT_HEAD, tax.TYPE_BEANIE, "Snowstar Patch", GENDER_GENERAL),
    "BEECH Original Patch Beanie": (tax.CAT_HEAD, tax.TYPE_BEANIE, "Original Patch", GENDER_GENERAL),
    "BEECH Snowstar Patch Beanie": (tax.CAT_HEAD, tax.TYPE_BEANIE, "Snowstar Patch", GENDER_GENERAL),
    "Cap-B445": (tax.CAT_HEAD, tax.TYPE_BEANIE, "Original Patch", GENDER_GENERAL),
    "Cap-B641": (tax.CAT_HEAD, tax.TYPE_CAP, "Patch Snapback Trucker", GENDER_GENERAL),
    "Westford Mill FairTrade Cotton Junior Apron": (
        tax.CAT_HOSP,
        tax.TYPE_APRON,
        "Fairtrade Cotton Junior",
        GENDER_KIDS,
    ),
    "WFMILL FairTrade Cotton Junior Apron": (
        tax.CAT_HOSP,
        tax.TYPE_APRON,
        "Fairtrade Cotton Junior",
        GENDER_KIDS,
    ),
    "AA77-Apron": (tax.CAT_HOSP, tax.TYPE_APRON, "Bib Apron", GENDER_GENERAL),
    "W364-Apron": (tax.CAT_HOSP, tax.TYPE_APRON, "Cotton Adult Apron", GENDER_GENERAL),
    "Sticker": (tax.CAT_STICKER, tax.TYPE_STICKER, tax.DEFAULT_STYLE, GENDER_GENERAL),
    "Stickers": (tax.CAT_STICKER, tax.TYPE_STICKER, "Circle", GENDER_GENERAL),
    "STICKER": (tax.CAT_STICKER, tax.TYPE_STICKER, "A4", GENDER_GENERAL),
    "Mug": (tax.CAT_MUGS, tax.TYPE_MUG, tax.DEFAULT_STYLE, GENDER_GENERAL),
    "Mug-M61": (tax.CAT_MUGS, tax.TYPE_MUG, "Ceramic Mug", GENDER_GENERAL),
    "Mask": (tax.CAT_ACC, tax.TYPE_MASK, tax.DEFAULT_STYLE, GENDER_GENERAL),
    "6M014V": (tax.CAT_ACC, tax.TYPE_MASK, tax.DEFAULT_STYLE, GENDER_GENERAL),
    "Badge": (tax.CAT_ACC, tax.TYPE_BADGE, "25 mm", GENDER_GENERAL),
    "Card": (tax.CAT_ACC, tax.TYPE_CARD, "A5", GENDER_GENERAL),
    "Photo Acrylic": (tax.CAT_ACC, tax.TYPE_PHOTO, tax.DEFAULT_STYLE, GENDER_GENERAL),
    "Baby Drawer Lock": (tax.CAT_ACC, tax.TYPE_LOCK, tax.DEFAULT_STYLE, GENDER_GENERAL),
    "Only-Design": (tax.CAT_IRON, tax.TYPE_IRON, tax.DEFAULT_STYLE, GENDER_GENERAL),
    "Kids-Tutu": (tax.CAT_BABY, tax.TYPE_TUTU, tax.DEFAULT_STYLE, GENDER_KIDS),
}

CL_LEFTOVER_RULES = CL_STANDARD_RULES
CL_STANDARD_RULES_FOLD = {k.casefold(): v for k, v in CL_STANDARD_RULES.items()}
CL_LEFTOVER_RULES_FOLD = CL_STANDARD_RULES_FOLD

SOURCE_BTC_UID = "btc_uid"
SOURCE_BTC_SPC = "btc_spc"
SOURCE_UNEEK_SHORT = "uneek_short"
SOURCE_UNEEK_CODE = "uneek_code"
SOURCE_CL_STANDARD = "cl_standard"
SOURCE_CL_GA = SOURCE_CL_STANDARD
SOURCE_CL_EXISTING = "cl_existing"
SOURCE_PLAIN_LEFTOVER = "plain_leftover"
SOURCE_NONE = ""

_KIDS_RE = re.compile(
    r"\b(children'?s|childrens?|kids?|child|youth|infant|bab(?:y|ies)|toddler|junior)\b",
    re.I,
)
_WOMENS_RE = re.compile(r"\b(women'?s|womens?|ladies|lady|girl)\b", re.I)
_PLAIN_KIDS_RE = re.compile(
    r"\b(children'?s|childrens?|kids?|youth|infant|bab(?:y|ies)|toddler|junior|girl'?s|girls|boy'?s|boys)\b",
    re.I,
)
_PLAIN_WOMENS_RE = re.compile(r"\b(women'?s|womens?|ladies|lady)\b", re.I)
_MENS_RE = re.compile(r"\b(men'?s|mens?|male|adults?)\b", re.I)
_UNISEX_RE = re.compile(r"\bunisex\b", re.I)
_HIVIS_RE = re.compile(r"hi[-\s]?vi[sz]", re.I)

_CL_BRAND_PREFIXES = (
    "fruit of the loom",
    "westford mill",
    "wfmill",
    "beechfield",
    "bagbase",
    "bagbas",
    "gildan",
    "fotl",
    "folt",
    "uneek",
    "yoko",
    "delta",
    "beech",
)

_NON_APPAREL = frozenset(
    {
        tax.CAT_BAGS,
        tax.CAT_IRON,
        tax.CAT_STICKER,
        tax.CAT_MUGS,
        tax.CAT_ACC,
        tax.CAT_HEAD,
        tax.CAT_HOSP,
    }
)

_GARMENT_PHRASES = tuple(
    sorted(
        (
            "full zip hooded sweatshirt",
            "full zip micro fleece jacket",
            "full zip microfleece jacket",
            "full zip fleece jacket",
            "full zip soft shell jacket",
            "full zip softshell jacket",
            "full zip fleece",
            "quarter zip microfleece jacket",
            "quarter zip sweatshirt",
            "1/4 zip micro fleece jacket",
            "hooded sweatshirt",
            "crewneck sweatshirt",
            "crew neck sweatshirt",
            "crew neck t-shirt",
            "crew neck t shirt",
            "long sleeve baseball t-shirt",
            "long sleeve poplin shirt",
            "short sleeve poplin shirt",
            "long sleeve t-shirt",
            "long sleeve t shirt",
            "hi viz short sleeve polo shirt",
            "hi-viz short sleeve polo shirt",
            "hi viz short sleeve t shirt",
            "hi-viz short sleeve t shirt",
            "hi-viz polo shirt",
            "hi viz polo shirt",
            "hi-vis class 2 waistcoat",
            "hi-vis baseball safety helmet",
            "childrens hi-viz waist coat",
            "printable softshell gilet",
            "printable soft shell jacket",
            "padded bodywarmer",
            "body warmer",
            "waist coat",
            "waistcoat",
            "athletic vest",
            "tank top",
            "polo shirt",
            "poloshirt",
            "t-shirt",
            "t shirt",
            "sweat jacket",
            "sweatshirt",
            "crewneck",
            "hoodie",
            "hooded",
            "softshell jacket",
            "soft shell jacket",
            "outdoor jacket",
            "fleece jacket",
            "active jacket",
            "bomber jacket",
            "road safety jacket",
            "safety helmet",
            "jogging pants",
            "jog bottoms",
            "cargo shorts",
            "cargo trousers",
            "cargo trouser",
            "trouser regular",
            "trouser long",
            "trouser short",
            "trousers",
            "trouser",
            "rugby shirt",
            "body suit",
            "bodysuit",
            "baseball t-shirt",
            "ringer t",
            "poplin full sleeve shirt",
            "poplin half sleeve shirt",
            "pinpoint oxford full sleeve shirt",
            "pinpoint oxford half sleeve shirt",
            "full sleeve shirt",
            "half sleeve shirt",
            "longsleeve poloshirt",
            "long sleeve",
            "short sleeve",
            "full zip",
            "quarter zip",
            "microfleece",
            "micro fleece",
            "softshell",
            "soft shell",
            "cardigan",
            "gilet",
            "jacket",
            "fleece",
            "beanie",
            "beanies",
            "snapback",
            "trucker cap",
            "dad cap",
            "panel cap",
            "apron",
            "tunic",
            "tabard",
            "scrub",
            "shorts",
            "polo",
            "vest",
            "shirt",
            "tote bag",
            "backpack",
            "shopper",
            "keyring",
            "helmet",
            "mask",
            "mug",
            "sticker",
            "stickers",
            "tutu",
            "bag",
        ),
        key=len,
        reverse=True,
    )
)

_KIND_STYLE = {
    "ironon": tax.DEFAULT_STYLE,
    "bag": tax.DEFAULT_STYLE,
    "cap": tax.DEFAULT_STYLE,
    "beanie": tax.DEFAULT_STYLE,
    "helmet": "Hi-Vis Baseball",
    "sticker": tax.DEFAULT_STYLE,
    "mug": tax.DEFAULT_STYLE,
    "mask": tax.DEFAULT_STYLE,
    "badge": "25 mm",
    "card": "A5",
    "photo": tax.DEFAULT_STYLE,
    "lock": tax.DEFAULT_STYLE,
    "tutu": tax.DEFAULT_STYLE,
    "apron": tax.DEFAULT_STYLE,
    "tee_hoodie": "Kids Set",
    "baby": "Baby Body Suit",
    "hi_viz_polo": "High Visibility",
    "hi_viz_tee": "High Visibility",
    "hi_viz_trouser": "High Visibility",
    "hi_viz_waistcoat": "High Visibility",
    "hi_viz_jacket": "High Visibility",
    "hi_vis": "High Visibility",
    "hoodie": tax.DEFAULT_STYLE,
    "sweatshirt": tax.DEFAULT_STYLE,
    "tee": tax.DEFAULT_STYLE,
    "long_sleeve": tax.DEFAULT_STYLE,
    "tank": tax.DEFAULT_STYLE,
    "vest": tax.DEFAULT_STYLE,
    "polo": tax.DEFAULT_STYLE,
    "long_polo": tax.DEFAULT_STYLE,
    "woven": tax.DEFAULT_STYLE,
    "trouser": tax.DEFAULT_STYLE,
    "shorts": tax.DEFAULT_STYLE,
    "jogger": tax.DEFAULT_STYLE,
    "jacket": tax.DEFAULT_STYLE,
    "fleece": tax.DEFAULT_STYLE,
    "gilet": tax.DEFAULT_STYLE,
    "cardigan": tax.DEFAULT_STYLE,
    "rugby": "Classic",
    "tunic": tax.DEFAULT_STYLE,
    "scrub": tax.DEFAULT_STYLE,
    "tabard": "Premium",
}

_CODE_RE = re.compile(r"[A-Za-z]{0,6}\d+[A-Za-z]{0,4}")
_IRON_SIZES = frozenset({"A3", "A4", "A5", "A6"})


def cell(value: object) -> str:
    if value is None:
        return ""
    s = str(value).strip()
    if s.lower() in ("nan", "none", "none"):
        return ""
    if s.endswith(".0") and s[:-2].isdigit():
        return s[:-2]
    return s


def keyfold(value: object) -> str:
    return cell(value).casefold()


def _norm_ga(ga: object) -> str:
    s = cell(ga)
    s = s.replace("\u2019", "'").replace("\u2018", "'")
    s = s.replace("TrouserLong", "Trouser Long")
    s = s.replace("Crew New", "Crewneck")
    s = re.sub(r"\s+", " ", s).strip()
    return s


def _fold_ga(ga: object) -> str:
    return _norm_ga(ga).casefold()


@dataclass(frozen=True)
class AreebValues:
    category: str = ""
    product_type: str = ""
    product_style: str = ""
    department: str = ""
    source: str = SOURCE_NONE

    def as_dict(self) -> dict[str, str]:
        return {
            "Category (Areeb)": self.category,
            "Product Type (Areeb)": self.product_type,
            "Product Style (Areeb)": self.product_style,
            "Department (Areeb)": self.department,
        }

    def any_filled(self) -> bool:
        return bool(self.category or self.product_type or self.product_style or self.department)

    def all_filled(self) -> bool:
        return bool(self.category and self.product_type and self.product_style and self.department)


def coalesce_areeb(base: AreebValues, extra: AreebValues) -> AreebValues:
    """Keep supplier cells; fill only the holes from leftover."""
    if not extra.any_filled() or base.all_filled():
        return base
    return AreebValues(
        category=base.category or extra.category,
        product_type=base.product_type or extra.product_type,
        product_style=base.product_style or extra.product_style,
        department=base.department or extra.department,
        source=base.source or extra.source,
    )


def from_src_row(row: Mapping[str, Any], src_map: dict[str, str], source: str) -> AreebValues:
    picked = {areeb: cell(row.get(src)) for areeb, src in src_map.items()}
    return AreebValues(
        category=picked["Category (Areeb)"],
        product_type=picked["Product Type (Areeb)"],
        product_style=picked["Product Style (Areeb)"],
        department=picked["Department (Areeb)"],
        source=source if any(picked.values()) else SOURCE_NONE,
    )


def _values_from_tuple(parts: tuple[str, str, str, str]) -> AreebValues:
    category, product_type, product_style, department = parts
    return AreebValues(
        category=category,
        product_type=product_type,
        product_style=product_style,
        department=department,
        source=SOURCE_CL_STANDARD if any(parts) else SOURCE_NONE,
    )


def _ironon_style(ga: str) -> str:
    lower = ga.casefold()
    for prefix in ("dtf-ironon-", "dtf-ironon", "ironon-", "ironon"):
        if lower.startswith(prefix):
            rest = ga[len(prefix) :].lstrip("-")
            return rest or "Iron-On"
    return "Iron-On"


def _strip_brand(folded: str) -> str:
    for prefix in _CL_BRAND_PREFIXES:
        if folded == prefix:
            return ""
        if folded.startswith(prefix + " ") or folded.startswith(prefix + "-"):
            return folded[len(prefix) :].lstrip(" -")
    return folded


def _title_style(text: str) -> str:
    bits: list[str] = []
    for raw in text.replace("-", " ").split():
        low = raw.casefold()
        if low in {"ux", "v"}:
            bits.append(raw.upper())
        elif low in {"hi-viz", "hi-vis", "hiviz", "hivis"}:
            bits.append("Hi-Vis")
        elif low in {"t", "shirt"} and bits and bits[-1] == "T":
            bits[-1] = "T-Shirt"
        else:
            bits.append(raw[:1].upper() + raw[1:].lower() if raw else raw)
    return " ".join(bits)


def _style_from_ga(ga: str, kind: str) -> str:
    orig = _norm_ga(ga)
    cf = _fold_ga(ga)
    if cf.startswith("dtf-ironon") or cf.startswith("ironon"):
        rest = _ironon_style(orig)
        if rest.upper() in _IRON_SIZES:
            return rest.upper()
        named = tax.style_from_code(rest) or tax.style_from_code(rest.replace("-", ""))
        if named:
            return named
        for token in _codes_in(rest):
            hit = tax.style_from_code(token)
            if hit:
                return hit
        if rest.casefold() in {"k-t", "kt"}:
            return "Kids T-Shirt"
        return tax.collapse_style(_title_style(rest) or tax.DEFAULT_STYLE)
    for token in _codes_in(orig):
        hit = tax.style_from_code(token)
        if hit:
            return hit
    for needle, canon in tax.STYLE_PHRASES:
        if needle in cf:
            return canon
    folded = _strip_brand(cf)
    folded = _KIDS_RE.sub(" ", folded)
    folded = _WOMENS_RE.sub(" ", folded)
    folded = _MENS_RE.sub(" ", folded)
    folded = _UNISEX_RE.sub(" ", folded)
    folded = _HIVIS_RE.sub(" ", folded)
    for phrase in _GARMENT_PHRASES:
        folded = folded.replace(phrase, " ")
    folded = re.sub(r"\s+", " ", folded).strip(" -")
    if folded.endswith(" t"):
        folded = folded[:-2].strip()
    leftover_map = {
        "heavy": "Heavy Cotton",
        "china": "China Bag",
    }
    if folded in leftover_map:
        return leftover_map[folded]
    if folded:
        titled = _title_style(folded)
        return tax.collapse_style(titled)
    return tax.collapse_style(_KIND_STYLE.get(kind, tax.DEFAULT_STYLE))


def _classify_ga_pattern(ga: str) -> AreebValues:
    hit = _garment_kind(ga)
    if not hit:
        return AreebValues()
    category, kind = hit
    dept = _department_for(ga, category)
    return _values_from_tuple(
        (category, _product_type(category, kind, ga), _style_from_ga(ga, kind), dept)
    )


def _snap_cl_areeb(values: AreebValues) -> AreebValues:
    """Hashim #038: warehouse fill picks from the closed list; do not invent."""
    if not values.any_filled():
        return values
    category, product_type, product_style, department = snap_warehouse(
        category=values.category,
        product_type=values.product_type,
        product_style=values.product_style,
        department=values.department,
    )
    return AreebValues(
        category=category,
        product_type=product_type,
        product_style=product_style,
        department=department or values.department,
        source=values.source,
    )


def _garment_kind(ga: str) -> tuple[str, str] | None:
    """Return (Category, kind) from Gender Apparel text. None = unknown."""
    cf = _fold_ga(ga)
    if not cf:
        return None
    if cf.startswith("dtf-ironon") or cf.startswith("ironon"):
        return tax.CAT_IRON, "ironon"
    if (
        cf.startswith("bg-")
        or cf.startswith("kc-bg")
        or cf.startswith("pc-")
        or cf.startswith("wb-")
        or cf.startswith("cc-w")
        or cf.startswith("eco ")
    ):
        return tax.CAT_BAGS, "bag"
    if cf.startswith("cap-"):
        return tax.CAT_HEAD, "cap"
    if any(
        token in cf
        for token in ("tote", "backpack", "shopper", "keyring", "china bag", "chinabag")
    ):
        return tax.CAT_BAGS, "bag"
    if _HIVIS_RE.search(cf):
        if "polo" in cf:
            return tax.CAT_POLO, "hi_viz_polo"
        if "t-shirt" in cf or "t shirt" in cf:
            return tax.CAT_TEE, "hi_viz_tee"
        if "trouser" in cf:
            return tax.CAT_SAFETY, "hi_viz_trouser"
        if "waist" in cf:
            return tax.CAT_SAFETY, "hi_viz_waistcoat"
        if "jacket" in cf or "bomber" in cf:
            return tax.CAT_SAFETY, "hi_viz_jacket"
        if "helmet" in cf:
            return tax.CAT_SAFETY, "helmet"
        return tax.CAT_SAFETY, "hi_vis"
    if "tutu" in cf:
        return tax.CAT_BABY, "tutu"
    if "waist coat" in cf or "waistcoat" in cf:
        return tax.CAT_SAFETY, "hi_viz_waistcoat"
    if "apron" in cf:
        return tax.CAT_HOSP, "apron"
    if "sticker" in cf:
        return tax.CAT_STICKER, "sticker"
    if cf == "mug" or cf.startswith("mug-"):
        return tax.CAT_MUGS, "mug"
    if "mask" in cf:
        return tax.CAT_ACC, "mask"
    if cf == "badge":
        return tax.CAT_ACC, "badge"
    if cf == "card":
        return tax.CAT_ACC, "card"
    if "photo acrylic" in cf:
        return tax.CAT_ACC, "photo"
    if "drawer lock" in cf:
        return tax.CAT_ACC, "lock"
    if cf == "only-design":
        return tax.CAT_IRON, "ironon"
    if any(
        token in cf
        for token in ("beanie", "snapback", "dad cap", "panel cap", "trucker cap", "pom pom")
    ):
        return tax.CAT_HEAD, "beanie" if "beanie" in cf else "cap"
    if "romper" in cf:
        return tax.CAT_BABY, "baby"
    if "body suit" in cf or "bodysuit" in cf or "toddler" in cf or cf.endswith("-bs"):
        return tax.CAT_BABY, "baby"
    if "tabard" in cf:
        return tax.CAT_HEALTH, "tabard"
    if "tunic" in cf:
        return tax.CAT_HEALTH, "tunic"
    if "scrub" in cf:
        return tax.CAT_HEALTH, "scrub"
    if "t-shirt" in cf and "hoodie" in cf:
        return tax.CAT_SETS, "tee_hoodie"
    if "trouser" in cf:
        return tax.CAT_TROUSER, "trouser"
    if re.search(r"\bshorts\b", cf) and "sleeve" not in cf:
        return tax.CAT_SHORTS, "shorts"
    if "jog" in cf:
        return tax.CAT_TROUSER, "jogger"
    if "rugby" in cf:
        return tax.CAT_SHIRT, "rugby"
    if "polo" in cf:
        if "longsleeve" in cf or "long sleeve" in cf:
            return tax.CAT_POLO, "long_polo"
        return tax.CAT_POLO, "polo"
    if "gilet" in cf or "bodywarmer" in cf or "body warmer" in cf:
        return tax.CAT_GILET, "gilet"
    if "cardigan" in cf:
        return tax.CAT_JACKET, "cardigan"
    if "fleece" in cf:
        return tax.CAT_FLEECE, "fleece"
    if any(
        token in cf
        for token in (
            "softshell",
            "soft shell",
            "outdoor jacket",
            "bomber",
            "padded",
            "active jacket",
        )
    ) or re.search(r"\bjacket\b", cf):
        if "sweat" in cf:
            return tax.CAT_SWEAT, "hoodie"
        return tax.CAT_JACKET, "jacket"
    if "hoodie" in cf or "hooded" in cf:
        return tax.CAT_SWEAT, "hoodie"
    if "sweatshirt" in cf or "crewneck" in cf:
        return tax.CAT_SWEAT, "sweatshirt"
    if "tank" in cf:
        return tax.CAT_TEE, "tank"
    if "vest" in cf:
        return tax.CAT_TEE, "vest"
    if "long sleeve" in cf and ("t-shirt" in cf or "t shirt" in cf or cf.endswith(" t")):
        return tax.CAT_TEE, "long_sleeve"
    if "t-shirt" in cf or "t shirt" in cf or re.search(r"\bt$", cf):
        return tax.CAT_TEE, "tee"
    if "poplin" in cf or "oxford" in cf or re.search(r"\bshirt\b", cf):
        return tax.CAT_SHIRT, "woven"
    return None


def _department_for(ga: str, category: str) -> str:
    if category == tax.CAT_BABY or category == tax.CAT_SETS:
        return GENDER_KIDS
    if category in _NON_APPAREL:
        if category == tax.CAT_HOSP and _KIDS_RE.search(ga):
            return GENDER_KIDS
        return GENDER_GENERAL
    if _KIDS_RE.search(ga):
        return GENDER_KIDS
    if _WOMENS_RE.search(ga):
        return GENDER_WOMENS
    if _UNISEX_RE.search(ga) or _HIVIS_RE.search(ga):
        return GENDER_UNISEX
    if category == tax.CAT_SAFETY or (category == tax.CAT_HEALTH and "scrub" in _fold_ga(ga)):
        return GENDER_UNISEX
    if _MENS_RE.search(ga):
        return GENDER_MENS
    if category in {
        tax.CAT_TEE,
        tax.CAT_SWEAT,
        tax.CAT_POLO,
        tax.CAT_SHIRT,
        tax.CAT_TROUSER,
        tax.CAT_SHORTS,
        tax.CAT_JACKET,
        tax.CAT_FLEECE,
        tax.CAT_GILET,
        tax.CAT_HEALTH,
    }:
        return GENDER_MENS
    return GENDER_GENERAL


def _codes_in(text: str) -> list[str]:
    found = _CODE_RE.findall(text or "")
    out: list[str] = []
    seen: set[str] = set()
    for raw in sorted(found, key=len, reverse=True):
        token = raw.upper()
        if token in _IRON_SIZES or token in seen:
            continue
        seen.add(token)
        out.append(token)
    return out


def _first_named_code(text: str) -> str:
    for token in _codes_in(text):
        hit = tax.style_from_code(token)
        if hit:
            return token
    return ""


def _bag_type(ga: str) -> str:
    code = _first_named_code(ga)
    hit = tax.bag_type_from_code(code)
    if hit:
        return hit
    cf = _fold_ga(ga)
    if "backpack" in cf:
        return tax.TYPE_BACKPACK
    if "book bag" in cf:
        return tax.TYPE_BOOK
    if "gymsac" in cf or "gym sac" in cf:
        return tax.TYPE_GYMSAC
    if "pencil" in cf:
        return tax.TYPE_PENCIL
    if "lunch" in cf or "cooler" in cf or "sandwich" in cf:
        return tax.TYPE_LUNCH
    if "drawstring" in cf:
        return tax.TYPE_DRAWSTRING
    if "barrel" in cf or "dance bag" in cf:
        return tax.TYPE_BARREL
    if any(tok in cf for tok in ("tote", "shopper", "bag for life", "china")):
        return tax.TYPE_TOTE
    return tax.TYPE_BAG_ACC


def _product_type(category: str, kind: str, ga: str) -> str:
    code = _first_named_code(ga)
    cf = _fold_ga(ga)
    by_kind = {
        "ironon": tax.TYPE_IRON,
        "bag": _bag_type(ga),
        "cap": tax.head_type_from_code(code) or tax.TYPE_CAP,
        "beanie": tax.head_type_from_code(code) or tax.TYPE_BEANIE,
        "helmet": tax.TYPE_HELMET,
        "sticker": tax.TYPE_STICKER,
        "mug": tax.TYPE_MUG,
        "mask": tax.TYPE_MASK,
        "badge": tax.TYPE_BADGE,
        "card": tax.TYPE_CARD,
        "photo": tax.TYPE_PHOTO,
        "lock": tax.TYPE_LOCK,
        "tutu": tax.TYPE_TUTU,
        "apron": tax.TYPE_APRON,
        "tee_hoodie": tax.TYPE_SET,
        "hi_viz_polo": tax.TYPE_HV_POLO,
        "hi_viz_tee": tax.TYPE_HV_TEE,
        "hi_viz_trouser": tax.TYPE_HV_TROUSER,
        "hi_viz_waistcoat": tax.TYPE_HV_VEST,
        "hi_viz_jacket": tax.TYPE_HV_JACKET,
        "hi_vis": tax.TYPE_HV_VEST,
        "tabard": tax.TYPE_TABARD,
        "tunic": tax.TYPE_TUNIC,
        "scrub": tax.TYPE_SCRUB,
        "jogger": tax.TYPE_JOGGER,
        "shorts": tax.TYPE_SHORTS,
        "rugby": tax.TYPE_RUGBY,
        "gilet": tax.TYPE_GILET,
        "cardigan": tax.TYPE_CARDIGAN,
        "fleece": tax.TYPE_FLEECE,
        "jacket": tax.TYPE_JACKET,
        "woven": tax.TYPE_SHIRT,
        "long_polo": tax.TYPE_LS_POLO,
        "polo": tax.TYPE_SS_POLO,
        "long_sleeve": tax.TYPE_LS_TEE,
        "tank": tax.TYPE_TANK,
        "vest": tax.TYPE_TANK,
        "tee": tax.TYPE_SS_TEE,
    }
    if kind == "baby":
        if "romper" in cf or tax.style_from_code(code) == "Baby Romper":
            return tax.TYPE_ROMPER
        if "body" in cf or tax.style_from_code(code) in {"Baby Body Suit", "Baby Bodysuit"}:
            return tax.TYPE_BODY
        return tax.TYPE_BABY_TEE
    if kind == "hoodie":
        if "full zip" in cf or "full-zip" in cf or "sweat jacket" in cf:
            return tax.TYPE_ZIP_HOODIE
        return tax.TYPE_HOODIE
    if kind == "sweatshirt":
        return tax.TYPE_SWEAT
    if kind == "trouser":
        if "cargo" in cf:
            return tax.TYPE_CARGO
        return tax.TYPE_TROUSER
    if kind in by_kind:
        return by_kind[kind]
    if category == tax.CAT_TEE:
        return tax.TYPE_SS_TEE
    if category == tax.CAT_SWEAT:
        return tax.TYPE_HOODIE
    if category == tax.CAT_POLO:
        return tax.TYPE_SS_POLO
    return category


def cl_standard(row: Mapping[str, Any]) -> AreebValues:
    """Warehouse Areeb 4-tuple from Gender Apparel. Department is gender only."""
    ga = _norm_ga(row.get("Gender Apparel"))
    if not ga:
        return AreebValues()
    exact = CL_STANDARD_RULES.get(ga) or CL_STANDARD_RULES_FOLD.get(ga.casefold())
    if exact:
        return _snap_cl_areeb(_values_from_tuple(exact))
    return _snap_cl_areeb(_classify_ga_pattern(ga))


leftover_cl = cl_standard
classify_cl = cl_standard


def _plain_gender(desc: str) -> str:
    if _PLAIN_KIDS_RE.search(desc):
        return GENDER_KIDS
    if _PLAIN_WOMENS_RE.search(desc):
        return GENDER_WOMENS
    if _UNISEX_RE.search(desc):
        return GENDER_UNISEX
    return GENDER_MENS


def _plain_pick(g: str, mens: str, ladies: str, kids: str, unisex: str | None = None) -> str:
    if g == GENDER_KIDS:
        return kids
    if g == GENDER_WOMENS:
        return ladies
    if g == GENDER_UNISEX:
        return unisex or mens
    return mens


def _plain_cat_type(desc: str) -> tuple[str, str]:
    """Map Plain Description → BTC Department / Sub Department. Empty desc → blank."""
    cf = _fold_ga(desc)
    if not cf:
        return "", ""
    g = _plain_gender(desc)
    if _HIVIS_RE.search(cf):
        ptype = "Childrens Hi-Vis" if g == GENDER_KIDS else "Hi Vis Clothing"
        return "Safetywear", ptype
    if any(tok in cf for tok in ("backpack", "tote", "shopper", "holdall", "duffle", "duffel", "rucksack")):
        return "Bags", "Bags, Backpacks Etc"
    if re.search(r"\bbag\b", cf) and "bodybag" not in cf:
        return "Bags", "Bags, Backpacks Etc"
    if any(tok in cf for tok in ("scarf", "neckwarmer", "snood")):
        return "HEADWEAR", "Scarves And Neckwear"
    if any(tok in cf for tok in ("beanie", "snapback", "trucker", "dad cap")) or re.search(
        r"\b(cap|hat|hats)\b", cf
    ):
        return "HEADWEAR", "Caps & Hats Etc"
    if "hoodie" in cf or "hooded" in cf or re.search(r"\bhoody\b", cf):
        return "SWEATSHIRTS AND HOODIES", _plain_pick(
            g,
            "Mens Sweatshirts & Hoodies",
            "Ladies Sweatshirts And Hoodies",
            "Childrens Sweatshirts And Hoodies",
            "UNISEX SWEATSHIRTS & HOODIES",
        )
    if "sweatshirt" in cf or "crewneck" in cf or "crew neck" in cf or re.search(r"\bsweat\b", cf):
        return "SWEATSHIRTS AND HOODIES", _plain_pick(
            g,
            "Mens Sweatshirts & Hoodies",
            "Ladies Sweatshirts And Hoodies",
            "Childrens Sweatshirts And Hoodies",
            "UNISEX SWEATSHIRTS & HOODIES",
        )
    if "polo" in cf:
        if "long sleeve" in cf or "longsleeve" in cf:
            return "POLO SHIRTS", _plain_pick(
                g,
                "Mens Long Sleeve Polo Shirts",
                "Ladies Long Sleeve Polo Shirts",
                "Childrens Polo Shirts",
                "Unisex Polo",
            )
        return "POLO SHIRTS", _plain_pick(
            g,
            "MENS SHORT SLEEVE POLO SHIRTS",
            "Ladies Short Sleeve Polo Shirts",
            "Childrens Polo Shirts",
            "Unisex Polo",
        )
    if "tank" in cf or re.search(r"\bvests?\b", cf):
        return "T-SHIRTS", _plain_pick(
            g,
            "Mens Tank Tops, Vest Etc",
            "Ladies Vests, Camisoles, Etc.",
            "Childrens T-Shirt",
            "UNISEX T-SHIRT",
        )
    if (
        "t-shirt" in cf
        or "t shirt" in cf
        or re.search(r"\btee\b", cf)
        or re.search(r"\bt$", cf)
        or " t/" in cf
        or cf.endswith(" t")
    ):
        if "long sleeve" in cf or "longsleeve" in cf or " lsl" in cf:
            return "T-SHIRTS", _plain_pick(
                g,
                "MENS LONG SLEEVE T-SHIRTS",
                "LADIES LONG SLEEVE T-SHIRTS",
                "Childrens T-Shirt",
                "Unisex Long Sleeve T-Shirt",
            )
        return "T-SHIRTS", _plain_pick(
            g,
            "MENS SHORT SLEEVE T-SHIRT",
            "Ladies Short Sleeve T-Shirts",
            "Childrens T-Shirt",
            "UNISEX T-SHIRT",
        )
    if "softshell" in cf or "soft shell" in cf:
        return "OUTERWEAR", _plain_pick(
            g,
            "Mens Softshell",
            "Ladies Softshell",
            "Childrens jackets",
            "Softshell",
        )
    if "fleece" in cf:
        return "OUTERWEAR", _plain_pick(
            g,
            "Mens Fleeces",
            "Ladies Fleeces",
            "Childrens Fleeces",
            "Unisex Fleece",
        )
    if any(tok in cf for tok in ("gilet", "bodywarmer", "body warmer", "body-warmer")):
        return "OUTERWEAR", _plain_pick(
            g,
            "MENS OUTER JACKETS",
            "Ladies Jackets",
            "Childrens jackets",
            "Unisex Jacket",
        )
    if any(
        tok in cf
        for tok in (
            "windbreaker",
            "jacket",
            "parka",
            "anorak",
            "raincoat",
            "shower",
            "bomber",
            "coat",
        )
    ):
        return "OUTERWEAR", _plain_pick(
            g,
            "MENS OUTER JACKETS",
            "Ladies Jackets",
            "Childrens jackets",
            "Unisex Jacket",
        )
    if any(tok in cf for tok in ("jogger", "jog pant", "jogpant", "sweatpant")) or re.search(
        r"\b(trousers?|pants)\b", cf
    ):
        return "TROUSERS & JOGPANTS", _plain_pick(
            g,
            "MENS TROUSERS AND JOGPANTS",
            "Ladies Trousers And Jog Pants",
            "Childrens Trousers And Jogpants",
            "UNISEX TROUSERS & JOGPANTS",
        )
    if re.search(r"\bshorts\b", cf) and "sleeve" not in cf:
        return "TROUSERS & JOGPANTS", _plain_pick(
            g,
            "Mens Shorts",
            "Ladies Shorts",
            "Childrens Trousers And Jogpants",
            "Mens Shorts",
        )
    if re.search(r"\bdress\b", cf):
        return "Corporate Wear", "Ladies Dress"
    if re.search(r"\bskirt\b", cf):
        return "Corporate Wear", "Ladies Skirts"
    if any(tok in cf for tok in ("apron", "tabard")):
        return "Hospitallity", "Aprons And Tabards"
    if "tunic" in cf:
        return "Hospitallity", "Ladies Tunic" if g != GENDER_MENS else "Aprons And Tabards"
    if "chef" in cf:
        return "Hospitallity", _plain_pick(
            g,
            "Unisex Chefswear",
            "Ladies Chefswear",
            "Unisex Chefswear",
            "Unisex Chefswear",
        )
    if "bib" in cf or "bodysuit" in cf or "body suit" in cf or "romper" in cf:
        return "Babywear", "Baby And Toddlerwear"
    if any(tok in cf for tok in ("knit", "jumper", "cardigan", "sweater")) and "sweatshirt" not in cf:
        return "Knitwear", _plain_pick(
            g,
            "Men's Knitwear",
            "Ladies Knitwear",
            "Men's Knitwear",
            "Men's Knitwear",
        )
    if "midlayer" in cf or "mid layer" in cf:
        return "OUTERWEAR", "Mens Midlayer"
    if any(tok in cf for tok in ("boot", "shoe", "footwear", "trainer")):
        return "Footwear", "Footwear"
    if "glove" in cf:
        return "Accessories", "Gloves"
    if "sock" in cf:
        return "Accessories", "Socks"
    if "poplin" in cf or "oxford" in cf or re.search(r"\bshirts?\b", cf):
        return "Shirts & Blouses", _plain_pick(
            g,
            "Mens Shirts",
            "Ladies Shirts",
            "Mens Shirts",
            "Mens Shirts",
        )
    if "wrap" in cf or "keyring" in cf or "key ring" in cf:
        return "Accessories", "Accessories"
    return "Accessories", "Accessories"


def plain_leftover(*, brand: object = "", description: object = "") -> AreebValues:
    """Fill leftover Plain Areeb from the row itself. Category/Type use BTC vocabulary."""
    desc = cell(description)
    brand_s = cell(brand)
    category, product_type = _plain_cat_type(desc)
    if not (category or product_type or brand_s or desc):
        return AreebValues()
    return AreebValues(
        category=category,
        product_type=product_type,
        product_style=brand_s,
        department=desc,
        source=SOURCE_PLAIN_LEFTOVER,
    )


class AreebCatalogs:
    """In-memory BTC + Uneek indexes (Plain / Packs only)."""

    def __init__(self) -> None:
        self.btc_by_uid: dict[str, dict[str, str]] = {}
        self.btc_by_spc: dict[str, dict[str, str]] = {}
        self.uneek_by_short: dict[str, dict[str, str]] = {}
        self.uneek_by_code: dict[str, dict[str, str]] = {}

    def btc_uid(self, sku: object) -> AreebValues:
        row = self.btc_by_uid.get(cell(sku)) or self.btc_by_uid.get(keyfold(sku))
        if not row:
            return AreebValues()
        return from_src_row(row, BTC_SRC, SOURCE_BTC_UID)

    def btc_spc(self, code: object) -> AreebValues:
        row = self.btc_by_spc.get(cell(code)) or self.btc_by_spc.get(keyfold(code))
        if not row:
            return AreebValues()
        return from_src_row(row, BTC_SRC, SOURCE_BTC_SPC)

    def uneek_short(self, code: object) -> AreebValues:
        k = cell(code)
        row = self.uneek_by_short.get(k) or self.uneek_by_short.get(k.casefold())
        if not row:
            return AreebValues()
        return from_src_row(row, UNEEK_SRC, SOURCE_UNEEK_SHORT)

    def uneek_product_code(self, code: object) -> AreebValues:
        k = cell(code)
        row = self.uneek_by_code.get(k) or self.uneek_by_code.get(k.casefold())
        if not row:
            return AreebValues()
        return from_src_row(row, UNEEK_SRC, SOURCE_UNEEK_CODE)

    def classify_plain(
        self,
        sku: object,
        product_code: object = "",
        brand: object = "",
        description: object = "",
    ) -> AreebValues:
        leftover = plain_leftover(brand=brand, description=description)
        hit = self.btc_uid(sku)
        if hit.any_filled():
            if hit.all_filled():
                return hit
            return coalesce_areeb(hit, leftover)
        hit = self.uneek_short(sku)
        if hit.any_filled():
            if hit.all_filled():
                return hit
            return coalesce_areeb(hit, leftover)
        hit = self.btc_spc(product_code)
        if hit.any_filled():
            if hit.all_filled():
                return hit
            return coalesce_areeb(hit, leftover)
        return leftover

    def classify_packs(
        self,
        *,
        item1_sku: object = "",
        product_code: object = "",
        channel_child_sku: object = "",
    ) -> AreebValues:
        hit = self.btc_uid(item1_sku)
        if hit.any_filled():
            return hit
        hit = self.btc_spc(product_code)
        if hit.any_filled():
            return hit
        for key in (channel_child_sku, item1_sku, product_code):
            hit = self.uneek_short(key)
            if hit.any_filled():
                return hit
            hit = self.uneek_product_code(key)
            if hit.any_filled():
                return hit
        return AreebValues()

    def classify_cl(self, row: Mapping[str, Any]) -> AreebValues:
        return cl_standard(row)


def _norm_row(row: Mapping[str, Any]) -> dict[str, str]:
    return {str(k): cell(v) for k, v in row.items()}


def load_btc_into(cat: AreebCatalogs, rows: list[Mapping[str, Any]]) -> None:
    for raw in rows:
        row = _norm_row(raw)
        uid = cell(row.get("UID") or row.get("Sku") or row.get("SKU"))
        if not uid or uid.startswith("["):
            continue
        if uid not in cat.btc_by_uid:
            cat.btc_by_uid[uid] = row
            cat.btc_by_uid[uid.casefold()] = row
        spc = cell(row.get("SPC"))
        if spc and spc not in cat.btc_by_spc:
            cat.btc_by_spc[spc] = row
            cat.btc_by_spc[spc.casefold()] = row


def load_uneek_into(cat: AreebCatalogs, rows: list[Mapping[str, Any]]) -> None:
    for raw in rows:
        row = _norm_row(raw)
        short = cell(row.get("Short Code"))
        if short and short not in cat.uneek_by_short:
            cat.uneek_by_short[short] = row
            cat.uneek_by_short[short.casefold()] = row
        code = cell(row.get("Product Code"))
        if code and code not in cat.uneek_by_code:
            cat.uneek_by_code[code] = row
            cat.uneek_by_code[code.casefold()] = row


def load_btc_csv(path: Path) -> list[dict[str, str]]:
    import csv

    last_err: Exception | None = None
    for enc in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            with path.open(encoding=enc, newline="") as f:
                return [{k: cell(v) for k, v in row.items()} for row in csv.DictReader(f)]
        except UnicodeDecodeError as exc:
            last_err = exc
            continue
    raise last_err or RuntimeError(f"Could not decode {path}")


def load_uneek_xlsx(path: Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows_iter = ws.iter_rows(values_only=True)
        headers = [cell(h) for h in next(rows_iter)]
        out: list[dict[str, str]] = []
        for raw in rows_iter:
            row = {headers[i]: cell(raw[i] if i < len(raw) else "") for i in range(len(headers)) if headers[i]}
            if any(row.values()):
                out.append(row)
        return out
    finally:
        wb.close()


def load_catalogs(
    *,
    btc_path: Path | None = None,
    uneek_path: Path | None = None,
) -> AreebCatalogs:
    from shared import paths as wh

    cat = AreebCatalogs()
    load_btc_into(cat, load_btc_csv(btc_path or wh.btc_product_data_path()))
    load_uneek_into(cat, load_uneek_xlsx(uneek_path or wh.uneek_product_data_path()))
    return cat


def apply_blank_only(current: Mapping[str, Any], values: AreebValues) -> dict[str, str]:
    """Return Areeb cells to write: only where current is blank and incoming is non-blank."""
    out: dict[str, str] = {}
    incoming = values.as_dict()
    for col in AREEB_COLS:
        if cell(current.get(col)):
            continue
        val = incoming[col]
        if val:
            out[col] = val
    return out


def apply_areeb(current: Mapping[str, Any], values: AreebValues) -> dict[str, str]:
    """Blank-only for supplier joins. CL standard overwrites all four Areeb cells.

    Hashim #038: an off-list warehouse style/type/category is written blank
    (do not keep an invented cell).
    """
    if values.source != SOURCE_CL_STANDARD:
        return apply_blank_only(current, values)
    out: dict[str, str] = {}
    for col, val in values.as_dict().items():
        new = cell(val)
        if cell(current.get(col)) != new:
            out[col] = new
    return out
