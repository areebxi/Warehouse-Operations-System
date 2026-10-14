"""BTC / Packs product lookup helpers for packing slips."""
from __future__ import annotations

import pandas as pd

from pdf_btc_images import _image_filename_from_url, _read_btc_product_data_csv
from pdf_constants import COLUMN_NAMES

def _load_btc_product_data_by_uid() -> dict[str, dict]:
    """Map BTC stock id (UID) to product fields when Database.xlsx has no row."""
    try:
        export_df, _enc = _read_btc_product_data_csv(dtype={"UID": str})
    except FileNotFoundError:
        return {}
    except Exception as e:
        print(f"[PDF] Warning: Could not load BTC Product Data fallback: {e}")
        return {}

    if "UID" not in export_df.columns:
        return {}

    export_df = export_df[~export_df["UID"].astype(str).str.startswith("[", na=False)]
    export_df = export_df.drop_duplicates(subset=["UID"], keep="first")

    by_uid: dict[str, dict] = {}
    for _, row in export_df.iterrows():
        uid = str(row.get("UID", "")).strip()
        if not uid:
            continue
        hi_res = row.get("image_url_high_res") or row.get("image_url_medium_res")
        by_uid[uid] = {
            COLUMN_NAMES["db_sku"]: uid,
            COLUMN_NAMES["product_code"]: str(row.get("SPC", "") or "").strip(),
            COLUMN_NAMES["brand"]: str(row.get("Brand", "") or "").strip(),
            COLUMN_NAMES["colour"]: str(row.get("Colour Name", "") or "").strip(),
            COLUMN_NAMES["size"]: str(row.get("Size", "") or "").strip(),
            COLUMN_NAMES["description"]: str(row.get("Description", "") or "").strip(),
            COLUMN_NAMES["package"]: "",
            COLUMN_NAMES["product_image_filename"]: _image_filename_from_url(hi_res),
            COLUMN_NAMES["brand_image_filename"]: _image_filename_from_url(row.get("brand image")),
        }
    return by_uid


def _row_to_product_dict(row) -> dict:
    if isinstance(row, pd.DataFrame):
        row = row.iloc[0]
    return row.to_dict()


def _lookup_product_details(
    products_df: pd.DataFrame,
    sku: str,
    export_by_uid: dict[str, dict] | None = None,
) -> dict:
    """Resolve product row from Database.xlsx, then BTC Product Data (UID = stock id)."""
    if not sku or (isinstance(sku, float) and pd.isna(sku)):
        return {}
    sku_str = str(sku).strip()
    if not sku_str:
        return {}

    try:
        return _row_to_product_dict(products_df.loc[sku_str])
    except KeyError:
        pass

    fallback = (export_by_uid or {}).get(sku_str)
    if fallback:
        print(f"     Using BTC Product Data fallback for SKU '{sku_str}' (not in Database.xlsx)")
        return dict(fallback)
    return {}


def _iter_pack_rows(path: str):
    from pathlib import Path

    p = Path(path)
    if p.suffix.lower() == ".csv":
        import csv

        with p.open(encoding="cp1252", newline="") as f:
            yield from csv.DictReader(f)
        return
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(it)]
        for raw in it:
            yield {
                headers[i]: raw[i]
                for i in range(min(len(headers), len(raw)))
                if headers[i]
            }
    finally:
        wb.close()


def _pack_cell_map(path: str, column: str) -> dict:
    out = {}
    for row in _iter_pack_rows(path):
        sku_str = str(row.get("Channel Child SKU") or "").strip()
        if not sku_str:
            continue
        normalized = sku_str.split("-", 1)[0] if "-" in sku_str else sku_str
        val = row.get(column)
        out[normalized] = "" if val is None else str(val).strip()
    return out


def _load_pack_names_map(excel_path: str) -> dict:
    """Load {normalized_pack_sku -> Pack Name}."""
    try:
        return _pack_cell_map(excel_path, "Pack Name")
    except Exception:
        return {}


def _load_pack_titles_map(excel_path: str) -> dict:
    """Load {normalized_pack_sku -> Supplier Name} (sheet column AI)."""
    try:
        return _pack_cell_map(excel_path, "Supplier Name")
    except Exception:
        return {}
