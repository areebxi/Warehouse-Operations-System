"""Packs Database loaders for PO stock / packing flows."""
from __future__ import annotations

import csv
import os
from pathlib import Path

from app_paths import packs_database_path

_COMPONENTS = tuple(range(1, 6))


def _pack_key(sku: str) -> str:
    sku = (sku or "").strip()
    return sku.split("-", 1)[0] if "-" in sku else sku


def basic_sku(original_sku: str) -> str:
    """Marketplace prefix: part before the first dash."""
    return _pack_key(original_sku)


def _iter_rows(path: str):
    p = Path(path)
    if p.suffix.lower() == ".csv":
        with p.open(encoding="cp1252", newline="") as f:
            yield from csv.DictReader(f)
        return
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.active
        it = ws.iter_rows(values_only=True)
        headers = [str(h).strip() if h is not None else "" for h in next(it)]
        for raw in it:
            yield {
                headers[i]: raw[i]
                for i in range(min(len(headers), len(raw)))
                if headers[i]
            }
    finally:
        wb.close()


def load_packs_database(excel_path=None):
    """Return {pack SKU -> [{sku, colour}, ...]} for Item 1–5."""
    packs_map = {}
    path = str(excel_path or packs_database_path())
    if not os.path.exists(path):
        print(f"[WARNING] Packs Database not found at: {path}")
        return packs_map
    try:
        for row in _iter_rows(path):
            pack_sku = str(row.get("Channel Child SKU") or "").strip()
            if not pack_sku:
                continue
            entries = []
            for n in _COMPONENTS:
                sku_str = str(row.get(f"Item {n} SKU") or "").strip()
                if not sku_str:
                    continue
                if "-" in sku_str:
                    sku_str = sku_str.split("-", 1)[0]
                colour = str(row.get(f"Item {n} Colour") or "").strip()
                entries.append({"sku": sku_str, "colour": colour})
            if entries:
                packs_map[_pack_key(pack_sku)] = entries
        print(f"[INFO] Packs Database loaded: {len(packs_map)} pack SKUs mapped")
        return packs_map
    except Exception as e:
        print(f"[WARNING] Failed to load Packs Database: {e}")
        return {}


def load_pack_names(excel_path=None):
    """Return {normalized pack SKU -> Pack Name}."""
    pack_names = {}
    path = str(excel_path or packs_database_path())
    if not os.path.exists(path):
        print(f"[WARNING] Packs Database not found at: {path}")
        return pack_names
    try:
        for row in _iter_rows(path):
            sku_str = str(row.get("Channel Child SKU") or "").strip()
            if not sku_str:
                continue
            name = row.get("Pack Name")
            pack_names[_pack_key(sku_str)] = "" if name is None else str(name).strip()
        print(f"[INFO] Pack Names loaded: {len(pack_names)} entries")
        return pack_names
    except Exception as e:
        print(f"[WARNING] Failed to load Pack Names: {e}")
        return {}
