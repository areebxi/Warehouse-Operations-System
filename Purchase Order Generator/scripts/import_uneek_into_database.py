"""
Append Uneek rows from Custom Label Database into Plain Database.xlsx.

SKU = CL Custom Label (Uneek Short Code). Existing SKUs are never modified.

  python scripts/import_uneek_into_database.py
  python scripts/import_uneek_into_database.py --dry-run
"""
from __future__ import annotations

import argparse
import csv
import shutil
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

import app_paths  # noqa: F401

from app_paths import PRODUCT_DATABASE_FILENAME, product_database_archive_dir, product_database_path

_ROOT = Path(__file__).resolve().parents[2]
if str(_ROOT) not in sys.path:
    sys.path.insert(0, str(_ROOT))

from shared import paths as wh  # noqa: E402

CORE_COLUMNS = [
    "SKU",
    "Product Code",
    "Brand",
    "Colour",
    "Size",
    "Description",
    "Product_Image_URL",
    "Brand_Image_URL",
    "Package",
]

_KNOWN_PACKAGES = {"Small Parcel", "Large Letter"}


def _cell(v: object) -> str:
    if v is None:
        return ""
    return str(v).strip()


def map_package(package_type: str) -> str:
    s = (package_type or "").strip()
    if s.upper().startswith("RM "):
        s = s[3:].strip()
    return s if s in _KNOWN_PACKAGES else ""


def description_from_ga(gender_apparel: str) -> str:
    s = (gender_apparel or "").strip()
    if s.casefold().startswith("uneek "):
        return s[6:].strip()
    return s


def product_image_filename(apparel_image: str) -> str:
    stem = (apparel_image or "").strip()
    if not stem:
        return ""
    return f"{stem}.jpg"


def uneek_to_plain_row(cl: dict[str, str]) -> dict[str, str]:
    return {
        "SKU": (cl.get("Custom Label") or "").strip(),
        "Product Code": (cl.get("Supplier Product Code") or "").strip(),
        "Brand": (cl.get("Brand") or "Uneek").strip() or "Uneek",
        "Colour": (cl.get("Colour") or "").strip(),
        "Size": (cl.get("Size") or "").strip(),
        "Description": description_from_ga(cl.get("Gender Apparel") or ""),
        "Product_Image_URL": product_image_filename(cl.get("Apparel Image") or ""),
        "Brand_Image_URL": "",
        "Package": map_package(cl.get("Package Type") or ""),
    }


def load_uneek_cl_rows(path: Path) -> list[dict[str, str]]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        return [
            r
            for r in csv.DictReader(f)
            if (r.get("Brand") or "").strip().casefold() == "uneek"
            and (r.get("Custom Label") or "").strip()
        ]


def backup_database(path: Path) -> Path:
    archive_dir = product_database_archive_dir()
    archive_dir.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    backup_path = archive_dir / f"{PRODUCT_DATABASE_FILENAME}.bak_{stamp}"
    shutil.copy2(path, backup_path)
    return backup_path


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--no-backup", action="store_true")
    ap.add_argument("--database", type=Path, default=product_database_path())
    args = ap.parse_args()

    db_path: Path = args.database
    if not db_path.is_file():
        print(f"Missing database: {db_path}", flush=True)
        return 1

    uneek = load_uneek_cl_rows(wh.cl_csv_path())
    print(f"Uneek CL rows: {len(uneek):,}", flush=True)

    wb = openpyxl.load_workbook(db_path)
    ws = wb.active
    header = [_cell(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
    if header[: len(CORE_COLUMNS)] != CORE_COLUMNS:
        print(f"Unexpected columns: {header}", flush=True)
        return 1

    existing: set[str] = set()
    for row in ws.iter_rows(min_row=2, max_col=1, values_only=True):
        sku = _cell(row[0])
        if sku:
            existing.add(sku.casefold())

    new_rows: list[dict[str, str]] = []
    skipped = 0
    seen: set[str] = set()
    for cl in uneek:
        row = uneek_to_plain_row(cl)
        key = row["SKU"].casefold()
        if not key or key in existing or key in seen:
            skipped += 1
            continue
        seen.add(key)
        new_rows.append(row)

    print(
        f"Existing SKUs: {len(existing):,}  append: {len(new_rows):,}  skip: {skipped:,}",
        flush=True,
    )
    if new_rows:
        sample = ", ".join(r["SKU"] for r in new_rows[:8])
        print(f"Sample new SKUs: {sample}", flush=True)

    assert map_package("RM Small Parcel") == "Small Parcel"
    assert map_package("RM Large Letter") == "Large Letter"
    assert description_from_ga("Uneek Ladies Shirt") == "Ladies Shirt"
    assert product_image_filename("UC711-WH-H") == "UC711-WH-H.jpg"

    if args.dry_run:
        print("Dry-run: no write.", flush=True)
        return 0
    if not new_rows:
        print("Nothing to add.", flush=True)
        return 0

    if not args.no_backup:
        backup = backup_database(db_path)
        print(f"Backup: {backup}", flush=True)

    for row in new_rows:
        ws.append([row[c] for c in CORE_COLUMNS])

    try:
        wb.save(db_path)
    except PermissionError:
        fallback = db_path.with_name(db_path.stem + "_uneek" + db_path.suffix)
        wb.save(fallback)
        print(f"Could not write {db_path} (file may be open). Wrote {fallback}", flush=True)
        return 0

    print(f"Updated: {db_path} (+{len(new_rows):,} Uneek rows)", flush=True)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
