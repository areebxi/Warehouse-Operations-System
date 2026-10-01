from __future__ import annotations
import pandas as pd
from fpdf import FPDF
from pathlib import Path
from datetime import date
from app_paths import APP_ROOT, asset_path, data_path, packs_database_path, product_database_path, tag_output_dir

def _load_pack_titles_map(excel_path: str) -> dict:
    """Load {normalized_pack_sku -> Title} using openpyxl (B=SKU, AI=Title)."""
    try:
        from openpyxl import load_workbook
    except Exception:
        return {}
    try:
        wb = load_workbook(excel_path, data_only=True, read_only=True)
        ws = wb.active
        pack_titles = {}
        # Column AI = 35 (1-based), so index 34 (0-based)
        for row in ws.iter_rows(min_row=2, max_col=35, values_only=True):
            pack_sku = row[1]  # Column B
            pack_title = row[34] if len(row) > 34 else None  # Column AI
            if not pack_sku:
                continue
            sku_str = str(pack_sku).strip()
            if not sku_str:
                continue
            normalized = sku_str.split('-')[0] if '-' in sku_str else sku_str
            pack_titles[normalized] = str(pack_title).strip() if pack_title is not None else ''
        return pack_titles
    except Exception:
        return {}
def _load_pack_names_map(excel_path: str) -> dict:
    """Load {normalized_pack_sku -> Pack Name} using openpyxl (B=SKU, C=Pack Name)."""
    try:
        from openpyxl import load_workbook
    except Exception:
        return {}
    try:
        wb = load_workbook(excel_path, data_only=True, read_only=True)
        ws = wb.active
        pack_names = {}
        for row in ws.iter_rows(min_row=2, max_col=3, values_only=True):
            pack_sku = row[1]
            pack_name = row[2] if len(row) > 2 else None
            if not pack_sku:
                continue
            sku_str = str(pack_sku).strip()
            if not sku_str:
                continue
            normalized = sku_str.split('-')[0] if '-' in sku_str else sku_str
            pack_names[normalized] = str(pack_name).strip() if pack_name is not None else ''
        return pack_names
    except Exception:
        return {}
def _lookup_product_details(
    products_df: pd.DataFrame,
    sku: str,
    export_by_uid: dict[str, dict] | None = None,
) -> dict:
    """Resolve product row from Database.xlsx, then BTC Product Data (UID = stock id)."""
    if not sku or (isinstance(sku, float) and pd.isna(sku)):
        return {}
    sku_str = str(sku).strip()
    if not sku_str:
        return {}

    try:
        return _row_to_product_dict(products_df.loc[sku_str])
    except KeyError:
        pass

    fallback = (export_by_uid or {}).get(sku_str)
    if fallback:
        print(f"     Using BTC Product Data fallback for SKU '{sku_str}' (not in Database.xlsx)")
        return dict(fallback)
    return {}
def _safe_add_single_page(pdf: 'PDF', item_data, product_data, item_count, total_items):
    """Render a single item page even if PDF.add_packing_slip doesn't exist."""
    add_single_fn = getattr(pdf, 'add_packing_slip', None)
    if callable(add_single_fn):
        add_single_fn(item_data, product_data, item_count, total_items)
        return
    # Manual render using internal helpers
    pdf.add_page(orientation='L')
    if hasattr(pdf, '_draw_header'):
        pdf._draw_header(item_data, product_data or {}, item_count, total_items)
    if hasattr(pdf, '_draw_product_details'):
        pdf._draw_product_details(item_data, product_data or {}, total_items)
def _read_btc_product_data_csv(**kwargs) -> tuple[pd.DataFrame, str]:
    """Load BTC Product Data trying common encodings (exports may be cp1252)."""
    last_err: Exception | None = None
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return pd.read_csv(BTC_PRODUCT_DATA_FILE, encoding=encoding, low_memory=False, **kwargs), encoding
        except UnicodeDecodeError as exc:
            last_err = exc
            continue
    if last_err is not None:
        raise last_err
    raise RuntimeError(f"Could not read BTC Product Data: {BTC_PRODUCT_DATA_FILE}")
def _row_to_product_dict(row) -> dict:
    if isinstance(row, pd.DataFrame):
        row = row.iloc[0]
    return row.to_dict()
